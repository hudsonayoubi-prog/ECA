"""Parity: Python build_monthly_workbook must match the prototype's monthlyWorkbook for every state and for national. Synthetic data only."""
import json, os, subprocess, sys, tempfile, unittest
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(ROOT, 'scripts'))
sys.path.insert(0, HERE)
import pp_export as X
import test_export as TE

GEN = '13 Oct 2026, 6:00 pm'


def run_js(fx):
    out = subprocess.run(['node', os.path.join(HERE, 'monthly_wb_js.js')], input=json.dumps(fx), capture_output=True, text=True)
    if out.returncode:
        raise RuntimeError(out.stderr)
    return json.loads(out.stdout)


def check(tc, fx, label):
    js = run_js(fx)
    d = dict(fx)
    if fx.get('snapshot'):
        d['snapshots'] = {fx['month']: fx['snapshot']}
    for st in [z['code'] for z in fx['states']] + [None]:
        py = TE.norm(json.loads(json.dumps(X.build_monthly_workbook(d, fx['month'], st, 'Test User', GEN))))
        a = js['national'] if st is None else js['state'][st]
        TE.compare(tc, a, py, '%s %s' % (label, st or 'national'))


class T(unittest.TestCase):
    def test_parity_closed_months(self):
        for seed in range(6):
            check(self, TE.fixture(seed), 'seed%d' % seed)

    def test_parity_month_to_date_and_other_months(self):
        for seed, today, month in ((20, '2026-09-17', '2026-09'), (21, '2026-10-13', '2026-10'), (22, '2026-10-13', '2026-08')):
            check(self, TE.fixture(seed, today=today, month=month), 'seed%d' % seed)

    def test_parity_sparse_and_empty(self):
        for seed in range(30, 34):
            check(self, TE.fixture(seed, n_visits=8), 'sparse%d' % seed)
        fx = TE.fixture(1); fx['visits'] = []
        check(self, fx, 'empty')

    def test_parity_finalised_month_uses_locked_snapshot(self):
        fx = TE.fixture(4)
        states = {p['id']: p['state'] for p in fx['people']}
        fx['snapshot'] = [{'id': 'sVIC0', 'name': 'Locked One', 'state': 'VIC', 'rank': 1, 'total': 91.25, 'qualified': True},
                          {'id': 'sNSW0', 'name': 'Locked Two', 'state': 'NSW', 'rank': 2, 'total': 80.04, 'qualified': True},
                          {'id': 'sVIC1', 'name': 'Locked Three', 'state': 'VIC', 'rank': 3, 'total': 70, 'qualified': True},
                          {'id': 'sVIC2', 'name': 'Unranked', 'state': 'VIC', 'rank': None, 'total': 0, 'qualified': False}]
        check(self, fx, 'final')
        d = {**fx, 'snapshots': {fx['month']: fx['snapshot']}}
        m = X.build_monthly_workbook(d, fx['month'], 'VIC', 'Test User', GEN)
        text = json.dumps(m['sheets'][0])
        self.assertIn('leaderboard final', text)
        self.assertIn('Leaderboard #1', text)
        self.assertIn('Locked One (91.3)', text)

    def test_command_line_summary(self):
        fx = TE.fixture(8)
        with tempfile.TemporaryDirectory() as td:
            dp = os.path.join(td, 'd.json')
            json.dump({**fx, 'by': 'Test'}, open(dp, 'w'))
            self.assertEqual(X.main([dp, fx['month'], '--summary', 'VIC,ALL', td]), 0)
            self.assertEqual(sorted(f for f in os.listdir(td) if f.endswith('.xlsx')),
                             ['ECA VIC monthly summary September 2026.xlsx', 'ECA national breakdown September 2026.xlsx'])
            self.assertEqual(X.main([dp, fx['month'], '--summary', 'VIC']), 2)

    def test_workbook_files(self):
        from openpyxl import load_workbook
        fx = TE.fixture(2)
        with tempfile.TemporaryDirectory() as td:
            for st, sheets in (('VIC', ['Summary', 'Share of shelf', 'Institutions', 'Student Marketeers', 'Audits']),
                               (None, ['Summary', 'By state', 'Share of shelf', 'Institutions', 'Student Marketeers', 'Audits'])):
                m = X.build_monthly_workbook(fx, fx['month'], st, 'Test User', GEN)
                path = os.path.join(td, m['filename'])
                X.write_xlsx(m, path)
                wb = load_workbook(path)
                self.assertEqual(wb.sheetnames, sheets)
                self.assertEqual(wb['Summary']['A1'].value.split(' · ')[0] in ('Victoria', 'National breakdown'), True)
                self.assertEqual(wb['Audits'].freeze_panes is not None, True)
            self.assertIn('ECA VIC monthly summary September 2026.xlsx', os.listdir(td))
            self.assertIn('ECA national breakdown September 2026.xlsx', os.listdir(td))
            ws = wb['By state']
            self.assertEqual(ws.max_row, 1 + 4 + 1)  # header, 4 states, National row (this sheet has no title lines)
            self.assertEqual(ws.cell(ws.max_row, 1).value, 'National')
            self.assertTrue(ws.cell(ws.max_row, 1).font.bold)


if __name__ == '__main__':
    unittest.main()
