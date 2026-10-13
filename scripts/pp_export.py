"""Monthly audit export (BR-54, BR-55): a Python port of `auditExportWorkbook` and `exportDashboard` in src/app.js,
plus a writer that produces the real .xlsx with native Excel charts.

Two steps, so each can be tested on its own:
  build_export(d, month, states, by, generated) -> model   pure; same sheets/rows/cells/charts as the prototype
  write_xlsx(model, path)                                  openpyxl; Dashboard with 9 native bar charts
tests/pp/test_export.py compares `build_export` with the prototype's own JavaScript, then opens the file.

Inputs `d` are the same records as pp_monthly.py, with visits carrying the extra audit-form fields (see ROW_FIELDS).
`generated` is the "13 Oct 2026, 6:00 pm" style stamp; the caller supplies it so the build stays pure.
Known gaps: dates use the same en-AU wording as Node (Sept, June, July); ids sort by plain string order.
"""
import math, re
from datetime import date

import pp_monthly as M
from pp_monthly import pct0, pct_or

LABELS = {'ed': 'Energy Drink', 'sf': 'Sugarfree', 'zero': 'Zero', 'editions': 'Editions', 'monster': 'Monster',
          'monsterUltra': 'Monster Ultra', 'v': 'V', 'vsf': 'V Sugarfree', 'vblue': 'V Blue', 'solo': 'Solo Energy'}
NAVY, TEAL, AMBER = '0B2A5B', '0F8B8D', 'B88A00'
SHORT_MONTH = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'June', 'July', 'Aug', 'Sept', 'Oct', 'Nov', 'Dec']
LONG_MONTH = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']


def month_label(m):
    return '%s %s' % (LONG_MONTH[int(m[5:7]) - 1], m[:4])


def fmt_date(s):
    if not s:
        return '—'
    d = M.parse(s)
    return '%d %s %d' % (d.day, SHORT_MONTH[d.month - 1], d.year)


def short_month_year(m):
    return '%s %s' % (SHORT_MONTH[int(m[5:7]) - 1], m[2:4])


def month_last_day(m):
    return M.add_days(M.shift_month(m, 1) + '-01', -1)


def file_safe(t):
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Za-z0-9 ._-]+', '', str(t))).strip()


def ynx(v):
    return 'Yes' if v in (1, '1') and v is not True else 'No' if v in (0, '0') and v is not False else ''


def money(v):
    try:
        return '$%.2f' % float(v)
    except (TypeError, ValueError):
        return '—'


def _present(x):
    return x is not None and x != ''


def guess_store_type(l):
    if l.get('type') == 'Vending':
        return 'Vending'
    if l.get('indNka') == 'NKA':
        return 'NKA'
    name = l.get('name', '')
    if re.search(r'(^|\s)(tavern|pub|hotel)\b|\bbar\b', name, re.I) and not re.search(r'(juice|espresso|coffee)', name, re.I):
        return 'Student Bar'
    return 'QSR'


def store_type_of(l):
    return None if not l else (l.get('storeType') or guess_store_type(l))


def comp_total(cf):
    return sum(M._num(n) for n in (cf or {}).values())


def shelf_share(rb, cf):
    t = M._num(rb) + comp_total(cf)
    return M._num(rb) / t if t else None


def comp_detail_text(cf, cd, is_v):
    out = []
    for c, n in (cf or {}).items():
        x = (cd or {}).get(c) or {}
        k = M._num(n)
        t = '%s: %s facing%s' % (c, _n(k), '' if k == 1 else 's')
        if _present(x.get('price')):
            t += ' · ' + money(x['price'])
        if not is_v and M._num(x.get('pos')):
            t += ' · %s POS' % _n(x['pos'])
        if not is_v and M._num(x.get('coolers')):
            t += ' · %s cooler%s' % (_n(x['coolers']), '' if M._num(x['coolers']) == 1 else 's')
        out.append(t)
    return out


def _n(v):
    f = float(v)
    return int(f) if f.is_integer() else f


def _num0(v):
    f = M._num(v)
    return _n(f)


def _res(v):
    r = dict(v.get('response') or {})
    r['ranging'] = v.get('ranging', r.get('ranging'))
    return r


def _round(x):
    return math.floor(x + 0.5)


def audits_for(d, locs_by, month, st):
    vs = [v for v in d['visits'] if str(v['date']).startswith(month) and (locs_by.get(v['loc']) or {}).get('state') == st]
    return sorted(vs, key=lambda v: (v['date'] + (v.get('time') or ''), str(v['id'])))


def build_export(d, month, states, by, generated):
    locs = M.derive_locs(d['locations'], d['unis'])
    locs_by = {l['id']: l for l in locs}
    live = [u for u in d['unis'] if u.get('status') != 'Removed']
    today = d['today']
    prods = [(k, qsr, vend, comp) for k, qsr, vend, comp in d.get('products', M.PRODUCTS)]
    ml = month_label(month)
    ready = month_last_day(month) <= today
    state_name = lambda st: next((x['name'] for x in d['states'] if x['code'] == st), st)
    counted_v = lambda v: v.get('status') == 'Completed' and M.approval_of(v) not in ('Rejected', 'Duplicate')

    head = ['Audit ID', 'Date', 'Time', 'Student Marketeer', 'Reports to', 'Visit status', 'Approval', 'Counted', 'University or TAFE', 'Type', 'Campus',
            'Location ID', 'Location', 'Location type', 'Store type', 'Contact name', 'Contact role',
            *['%s%s ranged' % (LABELS[k], ' (comp)' if comp else '') for k, _, _, comp in prods], 'Red Bull ranging %',
            'Red Bull price', 'Top competitor', 'Competitor price', 'Red Bull POS present', 'Red Bull cooler present', 'Bullring / Bull’s Eye', 'Channel', 'Classification',
            'Red Bull facings', 'Competitor facings', 'Share of shelf', 'Competitor POS seen', 'Competitor coolers seen', 'Competitor detail',
            'DPs added', 'DPs lost', 'Net DPs', 'Editions sold in', 'Edition flavours', 'POS placed', 'POS type', 'Coolers in', 'Coolers out',
            'Vending company', 'Machine ID', 'Issue', 'Notes', 'Next step', 'Follow-up', 'Submitted by', 'Submitted at', 'Approved by', 'Source']

    def row(v):
        l = locs_by.get(v['loc'])
        r = _res(v)
        rg = r.get('ranging') or {}
        p = next((x for x in d['people'] if x['name'] == v.get('staff')), None)
        counted = counted_v(v)
        rp = M.rb_range_pct(l, r.get('ranging'), prods)
        cf = v.get('compFacings')
        rbf = v.get('rbFacings')
        sh = shelf_share(rbf, cf) if rbf is not None else None
        sub = v.get('submittedAt')
        return [v['id'], v['date'], v.get('time') or '', v.get('staff') or '', (p or {}).get('reportsTo') or '', v.get('status'),
                M.approval_of(v) if v.get('status') == 'Completed' else '', 'Yes' if counted else 'No',
                (l or {}).get('uniName') or v.get('uni'), (l or {}).get('kind') or '', (l or {}).get('campus') or '', v['loc'],
                (l or {}).get('name') or '', (l or {}).get('type') or '', store_type_of(l) or '',
                v.get('contactName') or '', v.get('contactRole') or '',
                *[ynx(rg.get(k)) for k, _, _, _ in prods], '' if rp is None else pct0(rp),
                _num_or_blank(r.get('price')), r.get('comp') or '', _num_or_blank(r.get('compPrice')), ynx(r.get('pos')), ynx(r.get('cooler')),
                r.get('bull') or '', r.get('channel') or '', r.get('classification') or '',
                '' if rbf is None else rbf, _num0(comp_total(cf)) if cf is not None else '', pct0(sh) if sh is not None else '',
                '' if v.get('compPos') is None else v['compPos'], '' if v.get('compCoolers') is None else v['compCoolers'],
                '; '.join(comp_detail_text(cf, v.get('compDetail'), (l or {}).get('type') == 'Vending')) if cf is not None else '',
                _num0(v.get('dpAdded')), _num0(v.get('dpLost')), _num0(M._num(v.get('dpAdded')) - M._num(v.get('dpLost'))), _num0(v.get('editions')),
                v.get('flavours') or '', _num0(v.get('posPlaced')), v.get('posType') or '', _num0(v.get('coolersIn')), _num0(v.get('coolersOut')),
                r.get('vendCo') or '', r.get('machineId') or '', v.get('issue') or '', v.get('notes') or '', v.get('nextStep') or '', v.get('followUp') or '',
                v.get('submittedBy') or '', sub.replace('T', ' ')[:16] if sub else '', v.get('approvedBy') or '',
                'Demo' if v.get('demo') else 'Workbook import' if v.get('imported') else 'Live']

    # ---- Summary ----
    summary = []
    for st in states:
        allv = audits_for(d, locs_by, month, st)
        vs = [v for v in allv if counted_v(v)]
        sl = [l for l in locs if l.get('state') == st]
        e = M.exec_stats(d['visits'], sl, month + '-01', month_last_day(month) if ready else today, d.get('products', M.PRODUCTS))
        act = [l for l in sl if l['locStatus'] == 'Active']
        seen = {v['loc'] for v in vs}
        vis = sum(1 for l in act if l['id'] in seen)
        summary.append([state_name(st), len(allv), len(vs), sum(1 for v in allv if v.get('status') == 'Completed' and M.approval_of(v) == 'Pending'),
                        sum(1 for v in allv if M.approval_of(v) in ('Rejected', 'Duplicate')), len({v['staff'] for v in vs if v.get('staff')}), vis, len(act),
                        pct0(vis / len(act)) if act else '', _n(e['editions']), _n(e['pos']), _n(e['coolers']), _n(e['dpAdded'] - e['dpLost']),
                        pct_or(e['share'], ''), pct_or(e['rangingAudited'], ''), ', '.join(M.managers_for(d['people'], st))])
    tot = []
    for r in summary:
        for i, v in enumerate(r):
            cur = tot[i] if i < len(tot) else None
            new = ((cur or 0) + v) if isinstance(v, (int, float)) and not isinstance(v, bool) else cur
            if i < len(tot):
                tot[i] = new
            else:
                tot.append(new)
    sum_rows = [{'cells': r, 'bold': False} for r in summary]
    if len(states) > 1:
        sum_rows.append({'cells': ['All states', *tot[1:8], pct0(tot[6] / tot[7]) if tot[7] else '', *tot[9:13], '', '', ''], 'bold': True})
    sum_sheet = {'name': 'Summary', 'title': 'ECA audit export · %s%s' % (ml, ' · ' + states[0] if len(states) == 1 else ' · all states'),
                 'notes': [('Month closed on %s. Generated %s by %s.' % (fmt_date(month_last_day(month)), generated, by)) if ready else
                           'Month to date (%s). The full export generates on %s.' % (fmt_date(today), fmt_date(month_last_day(month))),
                           'One sheet per state lists every audit in the month, including scheduled, cancelled, rejected and duplicate ones. “Counted” = completed and not rejected or duplicate; the totals here use counted audits only.'],
                 'header': ['State', 'All records', 'Counted audits', 'Awaiting approval', 'Rejected / duplicate', 'Student Marketeers auditing', 'Locations visited',
                            'Active locations', 'Coverage', 'Editions sold in', 'POS placed', 'Coolers placed', 'Net DPs', 'Share of shelf', 'Red Bull ranging', 'Managers'],
                 'rows': sum_rows}

    def width(h):
        if re.search(r'Notes|Next step|Competitor detail|Issue', h):
            return 40
        if re.search(r'Location$|University|Student Marketeer|Reports to|Contact name', h):
            return 24
        return min(18, max(9, len(h) + 2))

    state_sheets = []
    for st in states:
        allv = audits_for(d, locs_by, month, st)
        state_sheets.append({'name': '%s audits' % st, 'title': '%s · audits · %s' % (state_name(st), ml),
                             'notes': ['%d record%s.%s' % (len(allv), '' if len(allv) == 1 else 's', ' Includes demo activity (prototype).' if any(v.get('demo') for v in allv) else '')],
                             'header': head, 'rows': [{'cells': row(v), 'bold': False} for v in allv], 'widths': [width(h) for h in head]})

    dash, data_sheet = _dashboard(d, locs, locs_by, live, month, states, ready, generated, state_name)
    single = len(states) == 1
    name = file_safe('ECA audit export %s %s%s.xlsx' % (month, states[0] if single else 'all states', '' if ready else ' month to date'))
    return {'filename': name, 'sheets': [dash, sum_sheet, *state_sheets, data_sheet]}


def _num_or_blank(v):
    return '' if not _present(v) else _n(float(v)) if float(v).is_integer() else float(v)


def _dashboard(d, locs, locs_by, live, month, states, ready, generated, state_name):
    single = len(states) == 1
    today = d['today']
    ml = month_label(month)
    prods = d.get('products', M.PRODUCTS)
    end = lambda m: month_last_day(m) if month_last_day(m) <= today else today
    locs_of = lambda sts: [l for l in locs if l.get('state') in sts]
    counted = lambda m, sts: [v for v in d['visits'] if str(v['date']).startswith(m) and (locs_by.get(v['loc']) or {}).get('state') in sts
                              and v.get('status') == 'Completed' and M.approval_of(v) not in ('Rejected', 'Duplicate')]

    def stats_for(m, ls, vs):
        act = [l for l in ls if l['locStatus'] == 'Active']
        ids = {l['id'] for l in ls}
        mine = [v for v in vs if v['loc'] in ids]
        seen = {v['loc'] for v in mine}
        e = M.exec_stats(d['visits'], ls, m + '-01', end(m), prods)
        vis = sum(1 for l in act if l['id'] in seen)
        return {'audits': len(mine), 'active': len(act), 'visited': vis, 'coverage': vis / len(act) if act else None,
                'sms': len({v['staff'] for v in mine if v.get('staff')}), 'editions': e['editions'], 'pos': e['pos'], 'coolers': e['coolers'],
                'share': e['share'], 'ranging': e['rangingAudited'], 'brands': e['brands'], 'rb': e['rbFacings'], 'comp': e['compFacings']}

    vs_now = counted(month, states)
    all_locs = locs_of(states)
    now = stats_for(month, all_locs, vs_now)
    prev_m = M.shift_month(month, -1)
    prev = stats_for(prev_m, all_locs, counted(prev_m, states))
    months6 = [dict(m=m, **stats_for(m, all_locs, counted(m, states))) for m in [M.shift_month(month, -i) for i in (5, 4, 3, 2, 1, 0)]]
    if single:
        groups = [{'key': u['code'], 'label': u.get('short') or u['name'], 'ls': [l for l in locs if l['uni'] == u['code']]} for u in live if u['state'] == states[0]]
    else:
        groups = [{'key': st, 'label': st, 'ls': locs_of([st])} for st in states]
    g_stats = [dict(g, **stats_for(month, g['ls'], vs_now)) for g in groups]
    top = lambda arr, k, n=10: sorted([x for x in arr if x[k] > 0], key=lambda x: -x[k])[:n]
    inst_all = [dict(label=(u.get('short') or u['name']) + ('' if single else ' (%s)' % u['state']),
                     **stats_for(month, [l for l in locs if l['uni'] == u['code']], vs_now)) for u in live if u['state'] in states]
    cnt = {}
    for v in vs_now:
        if v.get('staff'):
            cnt[v['staff']] = cnt.get(v['staff'], 0) + 1
    sm_rows = sorted([{'label': n + ('' if single else ' (%s)' % ((next((p for p in d['people'] if p['name'] == n), None) or {}).get('state') or '')), 'audits': c}
                      for n, c in cnt.items()], key=lambda x: -x['audits'])[:10]

    DATA = 'Chart data'
    data = []
    charts = []

    def ref(c, r1, r2):
        col = _col(c)
        return "'%s'!$%s$%d:$%s$%d" % (DATA, col, r1, col, r2)

    def block(title, header, rows, pct_cols=()):
        data.append([{'v': title, 's': 10}])
        data.append([{'v': h, 's': 1} for h in header])
        r1 = len(data) + 1
        use = rows if rows else [['No data yet'] + [0] * (len(header) - 1)]
        for r in use:
            data.append([{'v': v, 's': 2 if i in pct_cols and isinstance(v, (int, float)) else 0} for i, v in enumerate(r)])
        r2 = len(data)
        data.append([])
        return {'r1': r1, 'r2': r2, 'rows': use}

    def chart(title, b, series, direction='col', pct=False):
        charts.append({'title': title, 'dir': direction, 'pct': pct,
                       'cats': {'ref': ref(0, b['r1'], b['r2']), 'vals': [str(r[0]) for r in b['rows']]},
                       'series': [{'name': name, 'ref': ref(c, b['r1'], b['r2']), 'vals': [r[c] for r in b['rows']], 'color': color, 'pointColors': pcs}
                                  for name, c, color, pcs in [(s + (None,))[:4] if len(s) == 3 else s for s in series]]})

    gl = 'university & TAFE' if single else 'state'
    gdir = 'bar' if single else 'col'
    first = top(g_stats, 'audits', 10)
    g_pick = first + [g for g in g_stats if not g['audits'] and g['active']][:max(0, 6 - len(first))] if single else g_stats
    head0 = 'University / TAFE' if single else 'State'
    bA = block('Audits by %s' % gl, [head0, 'Audits completed'], [[g['label'], g['audits']] for g in g_pick])
    chart('Audits completed by %s' % gl, bA, [('Audits completed', 1, NAVY)], gdir)
    bC = block('Coverage by %s' % gl, [head0, 'Coverage'], [[g['label'], g['coverage'] if g['coverage'] is not None else 0] for g in g_pick if g['active']], (1,))
    chart('Coverage: share of active stores audited', bC, [('Coverage', 1, NAVY)], gdir, True)
    bE = block('Execution by %s' % gl, [head0, 'Editions sold in', 'POS placed', 'Coolers placed'], [[g['label'], g['editions'], g['pos'], g['coolers']] for g in g_pick])
    chart('Editions, POS and coolers by %s' % gl, bE, [('Editions sold in', 1, NAVY), ('POS placed', 2, TEAL), ('Coolers placed', 3, AMBER)], gdir)
    bS = block('Share of shelf by %s' % gl, [head0, 'Share of shelf'], [[g['label'], g['share']] for g in g_pick if g['share'] is not None], (1,))
    chart('Red Bull share of shelf by %s' % gl, bS, [('Share of shelf', 1, NAVY)], gdir, True)
    brand_tot = sum(now['brands'].values())
    brand_rows = [[b, n / brand_tot if brand_tot else 0] for b, n in sorted(now['brands'].items(), key=lambda kv: (0 if kv[0] == 'Red Bull' else 1, -kv[1]))]
    bB = block('Share of front facings by brand', ['Brand', 'Share of facings'], brand_rows, (1,))
    chart('Share of front facings: Red Bull vs competitors', bB, [('Share of facings', 1, '9AA3B2', [NAVY if r[0] == 'Red Bull' else None for r in bB['rows']])], 'bar', True)
    bT = block('Last 6 months', ['Month', 'Audits completed', 'Editions sold in', 'POS placed', 'Coolers placed', 'Coverage', 'Share of shelf'],
               [[short_month_year(x['m']), x['audits'], x['editions'], x['pos'], x['coolers'], x['coverage'] if x['coverage'] is not None else 0, x['share'] if x['share'] is not None else 0] for x in months6], (5, 6))
    chart('Audits completed: last 6 months', bT, [('Audits completed', 1, NAVY)])
    chart('Editions, POS and coolers: last 6 months', bT, [('Editions sold in', 2, NAVY), ('POS placed', 3, TEAL), ('Coolers placed', 4, AMBER)])
    bI = block('Top universities & TAFEs by audits', ['University / TAFE', 'Audits completed'], [[x['label'], x['audits']] for x in top(inst_all, 'audits')])
    chart('Top 10 universities & TAFEs by audits', bI, [('Audits completed', 1, NAVY)], 'bar')
    bM = block('Top Student Marketeers by audits', ['Student Marketeer', 'Audits completed'], [[x['label'], x['audits']] for x in sm_rows])
    chart('Top 10 Student Marketeers by audits', bM, [('Audits completed', 1, TEAL)], 'bar')

    # ---- dashboard sheet ----
    W = 16
    cells, merges, heights = {}, [], {}

    def put(r, c, v, s):
        cells.setdefault(r, {})[c] = {'v': v, 's': s}

    put(0, 0, 'ECA audit report · %s · %s' % (ml, (state_name(states[0]) if single else 'All states')), 3)
    heights[0] = 24
    demo = any(v.get('demo') for v in vs_now)
    put(1, 0, '%s · generated %s · Education Channel Audit Tool (ECA)%s' % ('Full month' if ready else 'Month to date at ' + fmt_date(today), generated,
                                                                             ' · prototype figures include demo activity' if demo else ''), 4)
    prev_l = month_label(prev_m).split(' ')[0][:3]

    def dl(a, b, pct=False):
        if a is None or b is None:
            return ''
        dd = _round((a - b) * 100) if pct else a - b
        dd = _n(dd)
        return 'Same as %s' % prev_l if dd == 0 else '%s %s%s vs %s' % ('▲' if dd > 0 else '▼', abs(dd), ' pts' if pct else '', prev_l)

    tiles = [('AUDITS COMPLETED', now['audits'], False, dl(now['audits'], prev['audits'])),
             ('COVERAGE OF ACTIVE STORES', now['coverage'], True, '%d of %d stores · %s' % (now['visited'], now['active'], dl(now['coverage'], prev['coverage'], True))),
             ('STUDENT MARKETEERS AUDITING', now['sms'], False, dl(now['sms'], prev['sms'])),
             ('EDITIONS SOLD IN', now['editions'], False, dl(now['editions'], prev['editions'])),
             ('POS PLACED', now['pos'], False, dl(now['pos'], prev['pos'])),
             ('COOLERS PLACED', now['coolers'], False, dl(now['coolers'], prev['coolers'])),
             ('SHARE OF SHELF', now['share'], True, 'No facings recorded' if now['share'] is None else 'Red Bull %s vs comp %s · %s' % (_n(now['rb']), _n(now['comp']), dl(now['share'], prev['share'], True))),
             ('RED BULL RANGING', now['ranging'], True, 'No ranging recorded' if now['ranging'] is None else 'Of SKUs at audited stores · %s' % dl(now['ranging'], prev['ranging'], True))]
    for t, (label, val, pct, note) in enumerate(tiles):
        c = t * 2
        put(3, c, label, 7); put(3, c + 1, '', 7)
        put(4, c, '—' if val is None else _n(val) if not pct else val, 9 if pct else 8); put(4, c + 1, '', 8)
        put(5, c, note, 11); put(5, c + 1, '', 11)
        for r in (3, 4, 5):
            merges.append('%s%d:%s%d' % (_col(c), r + 1, _col(c + 1), r + 1))
    heights[3], heights[4], heights[5] = 28, 36, 44
    hl = (M.state_monthly(d, month, states[0]) if single else M.national_monthly(d, month))['highlights']
    put(7, 0, 'Key points', 10)
    r = 8
    for t in (hl or ['No audits recorded this month.']):
        put(r, 0, '•  ' + t, 0)
        merges.append('A%d:%s%d' % (r + 1, _col(W - 1), r + 1))
        r += 1
    r += 1
    put(r, 0, 'Charts', 10)
    r += 2
    H = 18
    for i, ch in enumerate(charts):
        c0 = (i % 2) * 8
        r0 = r + (i // 2) * (H + 1)
        ch['from'] = {'c': c0, 'r': r0}
        ch['to'] = {'c': c0 + 8, 'r': r0 + H}
    last = r + math.ceil(len(charts) / 2) * (H + 1)
    put(last, 0, 'Figures count completed audits (rejected and duplicate audits excluded). Coverage = active stores audited at least once in the month. Share of shelf = Red Bull front facings ÷ all front facings at each store’s latest audit. Chart figures are on the Chart data sheet.', 4)
    merges.append('A%d:%s%d' % (last + 1, _col(W - 1), last + 1))
    grid = []
    for i in range(max(cells) + 1):
        row = cells.get(i, {})
        grid.append([row.get(c) for c in range(max(row) + 1)] if row else [])
    dash = {'name': 'Dashboard', 'raw': True, 'noGrid': True, 'cells': grid, 'merges': merges, 'heights': heights, 'widths': [11.5] * W, 'charts': charts}
    return dash, {'name': DATA, 'raw': True, 'cells': data, 'widths': [34, 18, 16, 16, 16, 14, 14]}


def _col(i):
    s = ''
    i += 1
    while i:
        m = (i - 1) % 26
        s = chr(65 + m) + s
        i = (i - 1) // 26
    return s


# ---------------------------------------------------------------- writer
def write_xlsx(model, path):
    """Write a sheet model (build_export or build_monthly_workbook) to .xlsx. A model without a Dashboard gets no charts."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.line import LineProperties
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, Font as DFont
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def F(**k):
        return Font(name='Arial', **k)
    navy_fill, grey_fill = PatternFill('solid', fgColor='FF0B2A5B'), PatternFill('solid', fgColor='FFF1F3F7')
    # 0 normal · 1 header · 2 percent · 3 title · 4 note · 5 bold · 6 bold percent · 7 KPI label · 8 KPI value · 9 KPI value % · 10 section · 11 KPI note
    STY = {0: dict(font=F(size=11)), 1: dict(font=F(size=11, bold=True, color='FFFFFFFF'), fill=navy_fill),
           2: dict(font=F(size=11), fmt='0%'), 3: dict(font=F(size=14, bold=True, color='FF0B2A5B')),
           4: dict(font=F(size=10, italic=True, color='FF6B7385')), 5: dict(font=F(size=11, bold=True)),
           6: dict(font=F(size=11, bold=True), fmt='0%'),
           7: dict(font=F(size=9, bold=True, color='FF6B7385'), fill=grey_fill, align=Alignment(vertical='bottom', wrap_text=True)),
           8: dict(font=F(size=22, bold=True, color='FF0B2A5B'), fill=grey_fill, fmt='#,##0', align=Alignment(horizontal='left', vertical='center')),
           9: dict(font=F(size=22, bold=True, color='FF0B2A5B'), fill=grey_fill, fmt='0%', align=Alignment(horizontal='left', vertical='center')),
           10: dict(font=F(size=12, bold=True, color='FF0B2A5B')),
           11: dict(font=F(size=10, italic=True, color='FF6B7385'), fill=grey_fill, align=Alignment(vertical='top', wrap_text=True))}

    def style(cell, s):
        o = STY[s or 0]
        cell.font = o['font']
        if 'fill' in o: cell.fill = o['fill']
        if 'fmt' in o: cell.number_format = o['fmt']
        if 'align' in o: cell.alignment = o['align']

    wb = Workbook()
    wb.remove(wb.active)
    names = []
    for sh in model['sheets']:
        n = re.sub(r'[\\/?*\[\]:]', ' ', str(sh['name']))[:31] or 'Sheet'
        k = 2
        while n in names:
            n = n[:28] + ' %d' % k
            k += 1
        names.append(n)
        sh['_title'] = n
    data_name = next((s['_title'] for s in model['sheets'] if s['name'] == 'Chart data'), None)
    for sh in model['sheets']:
        ws = wb.create_sheet(sh['_title'])
        if sh.get('raw'):
            if sh.get('noGrid'):
                ws.sheet_view.showGridLines = False
            for r, row in enumerate(sh['cells'], 1):
                for c, x in enumerate(row, 1):
                    if not x: continue
                    cell = ws.cell(r, c)
                    if x['v'] not in (None, ''): cell.value = x['v']
                    style(cell, x.get('s'))
            for m in sh.get('merges', []): ws.merge_cells(m)
            for r, h in (sh.get('heights') or {}).items(): ws.row_dimensions[int(r) + 1].height = h
            for i, w in enumerate(sh.get('widths', []), 1): ws.column_dimensions[get_column_letter(i)].width = w
            continue
        r = 1
        if sh.get('title'): ws.cell(r, 1, sh['title']); style(ws.cell(r, 1), 3); r += 1
        for n in sh.get('notes', []): ws.cell(r, 1, n); style(ws.cell(r, 1), 4); r += 1
        if sh.get('title') or sh.get('notes'): r += 1
        head_r = r
        for c, h in enumerate(sh['header'], 1):
            ws.cell(r, c, h); style(ws.cell(r, c), 1)
        r += 1
        pct = set(sh.get('pct', []))
        for row in sh['rows']:
            for c, v in enumerate(row['cells'], 1):
                cell = ws.cell(r, c)
                if v not in (None, ''): cell.value = v
                style(cell, (6 if row['bold'] else 2) if (c - 1) in pct else 5 if row['bold'] else 0)
            r += 1
        for i, h in enumerate(sh['header'], 1):
            w = (sh.get('widths') or [None] * 99)[i - 1] or min(48, max(10, len(str(h)) + 2, *[len(str(x['cells'][i - 1] if x['cells'][i - 1] is not None else '')) + 2 for x in sh['rows']] or [0]))
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(head_r + 1, 1)
        if sh['rows']: ws.auto_filter.ref = '%s%d:%s%d' % ('A', head_r, get_column_letter(len(sh['header'])), head_r + len(sh['rows']))

    def txt(sz, color, bold=False):
        cp = CharacterProperties(sz=sz, b=bold, solidFill=color, latin=DFont(typeface='Arial'))
        return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])

    dash = next((s for s in model['sheets'] if s['name'] == 'Dashboard'), None)
    if dash is None:
        wb.active = 0
        wb.save(path)
        return
    dash_ws = wb[dash['_title']]
    data_ws = wb[data_name]
    for ch in dash['charts']:
        bc = BarChart()
        bc.type = 'bar' if ch['dir'] == 'bar' else 'col'
        bc.grouping = 'clustered'
        bc.title = ch['title']
        bc.style = 2
        bc.gapWidth = 80 if len(ch['series']) > 1 else 55
        if len(ch['series']) > 1: bc.overlap = -5
        for s in ch['series']:
            m = re.match(r"'(.+)'!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", s['ref'])
            col_i = sum((ord(x) - 64) * 26 ** i for i, x in enumerate(reversed(m.group(2))))
            r1, r2 = int(m.group(3)), int(m.group(5))
            ref_v = Reference(data_ws, min_col=col_i, min_row=r1, max_row=r2)
            bc.add_data(ref_v, titles_from_data=False)
            ser = bc.series[-1]
            from openpyxl.chart.series import SeriesLabel
            ser.tx = SeriesLabel(v=s['name'])
            ser.graphicalProperties = GraphicalProperties(solidFill=s['color'])
            ser.invertIfNegative = False
            for j, pc in enumerate(s.get('pointColors') or []):
                if pc:
                    dp = DataPoint(idx=j)
                    dp.graphicalProperties = GraphicalProperties(solidFill=pc)
                    ser.dPt.append(dp)
            ser.dLbls = DataLabelList()
            ser.dLbls.showVal = True
            for a in ('showSerName', 'showCatName', 'showLegendKey', 'showPercent'): setattr(ser.dLbls, a, False)
            ser.dLbls.numFmt = '0%' if ch['pct'] else '#,##0'
            ser.dLbls.position = 'outEnd'
            ser.dLbls.txPr = txt(900, '374151', True)
        m = re.match(r"'(.+)'!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", ch['cats']['ref'])
        bc.set_categories(Reference(data_ws, min_col=1, min_row=int(m.group(3)), max_row=int(m.group(5))))
        bc.x_axis.delete = False
        bc.y_axis.delete = ch['dir'] == 'bar'
        bc.y_axis.scaling.min = 0
        allv = [float(v) for s in ch['series'] for v in s['vals'] if isinstance(v, (int, float))]
        if ch['pct'] or not allv or max(allv) == 0: bc.y_axis.scaling.max = 1
        bc.y_axis.number_format = '0%' if ch['pct'] else '#,##0'
        bc.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E5E8EE', w=6350))
        bc.x_axis.txPr = txt(900, '374151')
        bc.y_axis.txPr = txt(800, '6B7385')
        if ch['dir'] == 'bar': bc.x_axis.scaling.orientation = 'maxMin'
        bc.legend = None if len(ch['series']) == 1 else bc.legend
        if bc.legend is not None: bc.legend.position = 't'
        bc.anchor = TwoCellAnchor(editAs='oneCell', _from=AnchorMarker(col=ch['from']['c'], row=ch['from']['r']), to=AnchorMarker(col=ch['to']['c'], row=ch['to']['r']))
        dash_ws.add_chart(bc)
    wb.active = 0
    wb.save(path)


def main(argv=None):
    """python3 scripts/pp_export.py data.json 2026-09 VIC,NSW out_dir     audit export (states: comma list)
python3 scripts/pp_export.py data.json 2026-09 --summary VIC,ALL out_dir   monthly summary workbooks (ALL = national)"""
    import json, os, sys
    from datetime import datetime
    a = argv if argv is not None else sys.argv[1:]
    if len(a) < 4:
        print(main.__doc__)
        return 2
    data = json.load(open(a[0], encoding='utf-8'))
    month, states, out = a[1], a[2].split(','), a[3]
    summary = False
    if a[2] == '--summary':
        summary, states, out = True, a[3].split(','), a[4] if len(a) > 4 else None
        if out is None:
            print(main.__doc__)
            return 2
    data.setdefault('today', date.today().isoformat())
    stamp = datetime.now().strftime('%d %b %Y, %I:%M %p').lstrip('0').replace('AM', 'am').replace('PM', 'pm')
    os.makedirs(out, exist_ok=True)
    if summary:
        for st in states:
            model = build_monthly_workbook(data, month, None if st == 'ALL' else st, data.get('by', 'ECA flow'), stamp)
            write_xlsx(model, os.path.join(out, model['filename']))
            print('Wrote', os.path.join(out, model['filename']))
        return 0
    model = build_export(data, month, states, data.get('by', 'ECA flow'), stamp)
    path = os.path.join(out, model['filename'])
    write_xlsx(model, path)
    print('Wrote', path)
    return 0


if __name__ == '__main__':
    import sys
    sys.exit(main())


# ---------------------------------------------------------------- monthly summary workbooks (BR-53)
def _one(n):
    return '%.1f' % (math.floor(n * 10 + 0.5) / 10)


def build_monthly_workbook(d, month, st, by, generated):
    """The Excel download for a state summary (st = 'VIC') or the national breakdown (st = None). Same sheets as `monthlyWorkbook`."""
    locs = M.derive_locs(d['locations'], d['unis'])
    locs_by = {l['id']: l for l in locs}
    sm = M.state_monthly(d, month, st) if st else M.national_monthly(d, month)
    x = sm['x']
    ml = month_label(month)
    today = d['today']
    if month < today[:7]:
        status = 'Month closed · leaderboard final' if sm['final'] else 'Month closed · leaderboard in review'
    else:
        status = 'Month to date at %s' % fmt_date(today)
    notes = ['%s. Generated %s by %s.' % (status, generated, by),
             'Counts Completed audits; rejected and duplicate audits are excluded. Coverage = active locations visited at least once in the month.']
    R = lambda cells, bold=False: {'cells': cells, 'bold': bold}

    kpi = [('Audits completed', x['audits']), ('Active locations', x['active']), ('Locations visited', x['visited']), ('Coverage', x['coverage']),
           ('Student Marketeers auditing', x['auditors']), ('Distribution points added', x['dpAdded']), ('Distribution points lost', x['dpLost']),
           ('Net distribution points', x['net']), ('Editions sold in', x['editions']), ('POS placed', x['pos']), ('Coolers in', x['coolersIn']),
           ('Coolers out', x['coolersOut']), ('Share of shelf', pct_or(x['share'])), ('Red Bull front facings', x['rbFacings']),
           ('Competitor front facings', x['compFacings']), ('Red Bull ranging (audited locations)', pct_or(x['ranging'])),
           ('Competitor POS seen', x['compPos']), ('Competitor coolers seen', x['compCoolers']), ('Issues flagged', x['issues']), ('Awaiting approval', x['pending'])]

    def kpi_sheet(title, extra):
        rows = [R([k, (pct0(v) if x['active'] else '—') if k == 'Coverage' else _clean(v)]) for k, v in kpi] + [R(e) for e in extra]
        return {'name': 'Summary', 'title': title, 'notes': notes, 'header': ['Measure', 'Value'], 'rows': rows, 'widths': [34, 40]}

    def audit_sheet(vs):
        rows = []
        for v in sorted(vs, key=lambda v: v['date'] + (v.get('time') or '')):
            l = locs_by.get(v['loc'])
            cf, rbf = v.get('compFacings'), v.get('rbFacings')
            sh = shelf_share(rbf, cf) if rbf is not None else None
            rows.append(R([v['date'], v.get('staff') or '', (l or {}).get('state') or '', (l or {}).get('uniName') or v.get('uni'), (l or {}).get('campus') or '',
                           (l or {}).get('name') or v['loc'], (store_type_of(l) or '') if l else '', M.approval_of(v),
                           _num0(v.get('dpAdded')), _num0(v.get('dpLost')), _num0(v.get('editions')), _num0(v.get('posPlaced')), _num0(v.get('coolersIn')), _num0(v.get('coolersOut')),
                           '' if rbf is None else rbf, _num0(comp_total(cf)) if cf is not None else '', pct0(sh) if sh is not None else '',
                           '' if v.get('compPos') is None else v['compPos'], '' if v.get('compCoolers') is None else v['compCoolers'],
                           '; '.join(comp_detail_text(cf, v.get('compDetail'), (l or {}).get('type') == 'Vending')) if cf is not None else '',
                           v.get('contactName') or '', v.get('issue') or '', 'Demo' if v.get('demo') else 'Live']))
        return {'name': 'Audits', 'header': ['Date', 'Student Marketeer', 'State', 'University or TAFE', 'Campus', 'Location', 'Store type', 'Approval', 'DPs added', 'DPs lost',
                                             'Editions', 'POS', 'Coolers in', 'Coolers out', 'Red Bull facings', 'Competitor facings', 'Share of shelf', 'Competitor POS',
                                             'Competitor coolers', 'Competitor detail', 'Spoke to', 'Issue', 'Source'], 'rows': rows}

    def inst_sheet(insts):
        return {'name': 'Institutions', 'header': ['University or TAFE', 'Type', 'State', 'Active locations', 'Visited', 'Coverage', 'Audits', 'Net DPs', 'Editions', 'POS'], 'pct': [5],
                'rows': [R([i['u']['name'], i['u']['kind'], i['u']['state'], i['active'], i['visited'], i['coverage'] if i['active'] else None, i['audits'],
                            _n(i['net']), _n(i['editions']), _n(i['pos'])]) for i in insts]}

    def sm_sheet(sms):
        return {'name': 'Student Marketeers', 'header': ['Student Marketeer', 'State', 'Reports to', 'Audits', 'Locations', 'Net DPs', 'Editions', 'POS', 'Coolers in', 'Last audit'],
                'rows': [R([r['p']['name'], r['p']['state'], r['p'].get('reportsTo') or 'Not set', r['audits'], r['locations'], _n(r['net']), _n(r['editions']), _n(r['pos']),
                            _n(r['coolers']), r['last'] or '']) for r in sms]}

    tot_f = x['rbFacings'] + x['compFacings']
    brand_sheet = {'name': 'Share of shelf', 'header': ['Brand', 'Front facings', 'Share', 'Locations seen at'],
                   'rows': [R([b, _n(n), pct0(n / tot_f) if tot_f else '', x['shelfLocs'] if b == 'Red Bull' else x['ex']['present'].get(b, 0)])
                            for b, n in sorted(x['ex']['brands'].items(), key=lambda kv: (0 if kv[0] == 'Red Bull' else 1, -kv[1]))]}
    top_rows = [['Leaderboard #%d' % r['rank'], '%s (%s)' % (r['name'], _one(r['total']))] for r in sm['top']]
    if st:
        summary = kpi_sheet('%s · monthly summary · %s' % (sm['stName'], ml), top_rows + [['Sent to', ', '.join(sm['to']) or 'No managers set']])
        return {'filename': file_safe('ECA %s monthly summary %s.xlsx' % (st, ml)),
                'sheets': [summary, brand_sheet, inst_sheet(sm['insts']), sm_sheet(sm['sms']), audit_sheet(x['vs'])]}
    summary = kpi_sheet('National breakdown · %s' % ml, top_rows)
    rows = [R([r['st']['name'], r['insts'], r['y']['active'], r['y']['visited'], r['y']['coverage'] if r['y']['active'] else None, r['y']['audits'], r['sms'], r['activeSms'],
               _n(r['y']['net']), _n(r['y']['editions']), _n(r['y']['pos']), _n(r['y']['coolersIn']), pct_or(r['y']['share'], ''), pct_or(r['y']['ranging'], ''),
               r['y']['issues'], r['y']['pending'], ', '.join(r['managers'])]) for r in sm['states']]
    rows.append(R(['National', sum(r['insts'] for r in sm['states']), x['active'], x['visited'], x['coverage'] if x['active'] else None, x['audits'],
                   sum(r['sms'] for r in sm['states']), sum(r['activeSms'] for r in sm['states']), _n(x['net']), _n(x['editions']), _n(x['pos']), _n(x['coolersIn']),
                   pct_or(x['share'], ''), pct_or(x['ranging'], ''), x['issues'], x['pending'], ''], True))
    by_state = {'name': 'By state', 'header': ['State', 'Universities & TAFEs', 'Active locations', 'Visited', 'Coverage', 'Audits', 'Student Marketeers', 'SMs auditing', 'Net DPs',
                                               'Editions', 'POS', 'Coolers in', 'Share of shelf', 'Ranging', 'Issues', 'Awaiting approval', 'Managers'], 'pct': [4], 'rows': rows}
    all_insts = [i for z in d['states'] for i in M.state_monthly(d, month, z['code'])['insts']]
    all_sms = [r for z in d['states'] for r in M.state_monthly(d, month, z['code'])['sms']]
    return {'filename': file_safe('ECA national breakdown %s.xlsx' % ml),
            'sheets': [summary, by_state, brand_sheet, inst_sheet(all_insts), sm_sheet(all_sms), audit_sheet(x['vs'])]}


def _clean(v):
    return _n(v) if isinstance(v, float) and v.is_integer() else v
