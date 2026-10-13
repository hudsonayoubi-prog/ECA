"""Turn the Dataverse build pack (xlsx) into a machine-readable schema.

Reads  docs/power-platform/ECA_Dataverse_Build_Pack.xlsx  ("Tables & columns" and "Choices" sheets)
Writes docs/power-platform/schema/eca-schema.json

Usage (repo root):  python3 scripts/pp_schema.py [pack.xlsx] [out.json]
The JSON is the input to scripts/pp_deploy.py. The workbook stays the source of truth.
"""
import json, os, re, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PACK = os.path.join(ROOT, 'docs/power-platform/ECA_Dataverse_Build_Pack.xlsx')
OUT = os.path.join(ROOT, 'docs/power-platform/schema/eca-schema.json')
PREFIX = 'eca_'
# Choice sets whose columns say "Choice (as X)" or whose inline label list does not match a set exactly.
CHOICE_ALIAS = {'Choice (as Results Distribution)': 'Distribution Method'}
INLINE_NAMES = {'Active / Removed': 'Institution Status', 'Pending / Approved / Rejected': 'Removal Status'}


def slug(s):
    return re.sub(r'[^a-z0-9]+', '', s.lower())


def load_choices(wb):
    sets = {}
    for name, label, order in list(wb['Choices'].iter_rows(min_row=2, values_only=True)):
        if name:
            sets.setdefault(name, []).append({'label': label, 'order': order})
    return sets


def find_choice_set(spec, sets):
    """spec is the text after 'Choice:'. Either a set name or 'A / B / C' labels."""
    spec = spec.strip()
    if spec in sets:
        return spec
    spec = re.sub(r'\s*\(blank = not checked\)', '', spec)
    labels = [x.strip() for x in spec.split(' / ')]
    for n, items in sets.items():
        if [i['label'] for i in items] == labels:
            return n
    # Not on the Choices sheet: create a set from the inline labels (e.g. Institution Status).
    name = INLINE_NAMES.get(' / '.join(labels), ' / '.join(labels))
    sets[name] = [{'label': l, 'order': i + 1} for i, l in enumerate(labels)]
    return name


def parse_default(text):
    m = re.search(r'Default (\w+)', text or '')
    if not m:
        return None
    v = m.group(1)
    return True if v == 'Yes' else False if v == 'No' else int(v) if v.isdigit() else v


def parse_type(dt, sets):
    """Return (kind, props) for a data-type string from the sheet."""
    primary = '· Primary name' in dt
    dt = dt.replace('· Primary name', '').strip()
    p = {}
    if dt.startswith('Lookup'):
        target = dt.split('→')[1].strip()
        return 'lookup', {'target': 'systemuser' if target.startswith('User') else target}
    if dt.startswith('Choice'):
        spec = dt.split(':', 1)[1] if ':' in dt else dt
        spec = CHOICE_ALIAS.get(dt, spec)
        return 'choice', {'choiceSet': find_choice_set(spec, sets)}
    if dt.startswith('Autonumber'):
        m = re.search(r'Autonumber (\S+)', dt)
        return 'autonumber', {'format': m.group(1)} if m else {'format': None}
    m = re.match(r'(Text|Email|Phone) \((\d+)\)', dt)
    if m:
        fmt = {'Text': 'Text', 'Email': 'Email', 'Phone': 'Phone'}[m.group(1)]
        return 'string', {'maxLength': int(m.group(2)), 'format': fmt}
    m = re.match(r'Multiline text(?: \((\d+)\))?', dt)
    if m:
        return 'memo', {'maxLength': int(m.group(1) or 100000)}
    if dt.startswith('Whole number'):
        lo, hi = None, None
        if '≥ 0' in dt: lo = 0
        elif '≥ 1' in dt: lo = 1
        m = re.search(r'(\d+)[–-](\d+)', dt)
        if m: lo, hi = int(m.group(1)), int(m.group(2))
        return 'integer', {'min': lo, 'max': hi}
    if dt.startswith('Decimal'):
        m = re.search(r'(\d+)[–-](\d+)', dt)
        return 'decimal', {'precision': 1 if '1 dp' in dt else 2,
                           'min': int(m.group(1)) if m else None, 'max': int(m.group(2)) if m else None}
    if dt == 'Currency': return 'money', {}
    if dt == 'Yes/No': return 'boolean', {}
    if dt == 'Date only': return 'date', {}
    if dt == 'Date and time': return 'datetime', {}
    if dt == 'Image': return 'image', {}
    if dt.startswith('File'): return 'file', {}
    raise ValueError('Unrecognised data type %r' % dt)


def build(pack=PACK):
    wb = openpyxl.load_workbook(pack)
    sets = load_choices(wb)
    tables, cur = [], None
    for tname, tlog, cname, clog, dt, req, dflt, notes in wb['Tables & columns'].iter_rows(min_row=2, values_only=True):
        if not tname:
            continue
        if cname == '(table)':
            cur = {'displayName': tname, 'logicalName': tlog, 'description': notes, 'columns': []}
            tables.append(cur)
            continue
        primary = '· Primary name' in (dt or '')
        # Visit ID is "Autonumber or text": imports need free text, so create it as text (see docs).
        if dt and dt.startswith('Autonumber or text'):
            dt = 'Text (50) · Primary name'
        kind, props = parse_type(dt, sets)
        col = {'displayName': cname, 'logicalName': clog, 'kind': kind, 'required': req == 'Required',
               'primary': primary, 'default': parse_default(dflt), 'behaviour': dflt, 'notes': notes, **props}
        cur['columns'].append(col)
    # Alternate keys named in the build pack notes.
    keys = {'eca_location': ['eca_name'], 'eca_leaderboardsnapshot': ['eca_name'], 'eca_campus': ['eca_name'],
            'eca_institution': ['eca_name'], 'eca_state': ['eca_name']}
    for t in tables:
        t['alternateKey'] = keys.get(t['logicalName'])
    used = {c['choiceSet'] for t in tables for c in t['columns'] if c['kind'] == 'choice'}
    chs = [{'name': n, 'logicalName': PREFIX + slug(n), 'options': [{'label': i['label'], 'value': 100000000 + i['order'] - 1} for i in items]}
           for n, items in sets.items() if n in used]
    schema = {'publisherPrefix': PREFIX, 'solution': 'ECA', 'choices': chs, 'tables': tables}
    check(schema)
    return schema


def check(schema):
    """Fail loudly if the pack is internally inconsistent."""
    names = {t['displayName']: t for t in schema['tables']}
    chs = {c['name'] for c in schema['choices']}
    for t in schema['tables']:
        prim = [c for c in t['columns'] if c['primary']]
        assert len(prim) == 1 and prim[0]['logicalName'] == PREFIX + 'name', 'primary name: ' + t['displayName']
        seen = set()
        for c in t['columns']:
            assert c['logicalName'] not in seen, 'duplicate column %s.%s' % (t['logicalName'], c['logicalName'])
            seen.add(c['logicalName'])
            if c['kind'] == 'lookup':
                assert c['target'] == 'systemuser' or c['target'] in names, 'lookup target ' + c['target']
            if c['kind'] == 'choice':
                assert c['choiceSet'] in chs


def main():
    pack = sys.argv[1] if len(sys.argv) > 1 else PACK
    out = sys.argv[2] if len(sys.argv) > 2 else OUT
    schema = build(pack)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(schema, f, ensure_ascii=False, indent=1)
        f.write('\n')
    ncol = sum(len(t['columns']) for t in schema['tables'])
    print('Wrote %s: %d tables, %d columns, %d choice sets' % (out, len(schema['tables']), ncol, len(schema['choices'])))


if __name__ == '__main__':
    main()
