import sys, re, datetime as dt, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

p = sys.argv[1]
wb = openpyxl.load_workbook(p)
TODAY = dt.datetime(2026, 10, 8)

NAVY = '1F3864'; GREEN = '375623'
thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
BF = Font(name='Arial', size=10)
AUTO = PatternFill('solid', fgColor='EDEDED')
NOTE = Font(name='Arial', italic=True, size=9, color='666666')

# name, full name, ID prefix, table stem, named-range stem, master label
UNIS = [
    ('Monash', 'Monash University', 'MON', 'Monash', 'Monash', 'Monash'),
    ('La Trobe', 'La Trobe University', 'LTU', 'LaTrobe', 'LaTrobe', 'La Trobe'),
    ('Deakin', 'Deakin University', 'DEA', 'Deakin', 'Deakin', 'Deakin'),
    ('Victoria University', 'Victoria University', 'VU', 'Victoria', 'Victoria', 'Victoria University'),
    ('RMIT', 'RMIT University', 'RMIT', 'RMIT', 'RMIT', 'RMIT'),
    ('Swinburne', 'Swinburne University', 'SWI', 'Swinburne', 'Swinburne', 'Swinburne'),
    ('Melbourne University', 'University of Melbourne', 'UOM', 'Melbourne', 'Melbourne', 'Melbourne University'),
    ('UTAS', 'University of Tasmania', 'UTAS', 'UTAS', 'UTAS', None),
]
for u in UNIS:
    if u[0] in wb.sheetnames: del wb[u[0]]
for n in ['Staff Management', 'Progress']:
    if n in wb.sheetnames: del wb[n]
for n in list(wb.defined_names.keys()):
    del wb.defined_names[n]

def q(sheet):
    return f"'{sheet}'" if re.search(r"[^A-Za-z0-9_]", sheet) else sheet

def table(ws, name, ref, style='TableStyleMedium2'):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)

def dv(ws, rngs, **kw):
    d = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
    ws.add_data_validation(d)
    for r in rngs: d.add(r)

def head(ws, row, col, text, fill=NAVY):
    c = ws.cell(row, col, text)
    c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    c.fill = PatternFill('solid', fgColor=fill)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    return c

def parse_date(v):
    if isinstance(v, dt.datetime): return v
    if isinstance(v, str):
        try: return dt.datetime.strptime(v.strip(), '%d/%m/%Y')
        except ValueError: return None
    return None

# ------------------------------------------------------------------ gather accounts from VIC Master
M = wb['VIC Master']; MH = 17
mh = {M.cell(MH, c).value: c for c in range(1, M.max_column + 1)}
mfirst = MH + 1
mlast = max(r for r in range(mfirst, M.max_row + 1) if M.cell(r, 2).value)
REF = re.compile(r"=IF\('?(.+?)'?!([A-Z]+\d+)=")
def src_val(r, col):
    m = REF.match(str(M.cell(r, mh[col]).value or ''))
    return wb[m.group(1)][m.group(2)].value if m else None

def clean_campus(c):
    c = str(c or '').strip()
    c = {'Monash Clayton': 'Clayton', 'Caulfield Monash': 'Caulfield', 'Caufield': 'Caulfield'}.get(c, c)
    return c

accounts = {u[5]: [] for u in UNIS if u[5]}
for r in range(mfirst, mlast + 1):
    lab = M.cell(r, 2).value
    name = str(src_val(r, 'Account') or '').strip()
    street = str(src_val(r, 'Street') or '').strip()
    city = str(src_val(r, 'City') or '').strip()
    addr = street if (not city or city.lower() in street.lower()) else f'{street}, {city}'
    accounts[lab].append(dict(row=r, name=name, campus=clean_campus(src_val(r, 'Campus')), addr=addr.strip(', '),
                              last=parse_date(src_val(r, 'Last Visit')), nxt=str(src_val(r, 'Next Step') or '').strip()))

# ------------------------------------------------------------------ Staff Management
S = wb.create_sheet('Staff Management')
S.sheet_properties.tabColor = '7030A0'
heads = ['Staff Member', 'University', 'Campus', 'Status']
for i, h in enumerate(heads, 1): S.cell(1, i, h)
staff = [('George Johnson', 'Monash University', 'Clayton'), ('George Johnson', 'Monash University', 'Caulfield'),
         ('Mill Burgess', 'Victoria University', 'All'), ('Alex Mozjerin', 'RMIT University', 'All'),
         ('Kristie Phan', 'Swinburne University', 'All'), ('Ella Dzielakowski', 'Swinburne University', 'All'),
         ('Abi James', 'University of Tasmania', 'All'), ('Lachlan Heath', 'University of Tasmania', 'All'),
         ('Amanda Shen', 'University of Melbourne', 'All'), ('Andy Sitch', 'University of Melbourne', 'All')]
people = []
for s_, _, _ in staff:
    if s_ not in people: people.append(s_)
for uni in ['La Trobe University', 'Deakin University']:
    for s_ in people: staff.append((s_, uni, 'All'))
SROWS = 80  # table rows 2..81
for i, (a, b, c) in enumerate(staff, 2):
    S.cell(i, 1, a); S.cell(i, 2, b); S.cell(i, 3, c); S.cell(i, 4, 'Active')
send = SROWS + 1
table(S, 'tblStaffAllocations', f'A1:D{send}', 'TableStyleMedium4')
S['F1'] = 'Helper: first row for staff + uni'; S['G1'] = 'Helper: rank within uni'; S['H1'] = 'Helper: key'
for r in range(2, send + 1):
    S.cell(r, 6, f'=IF(A{r}="","",IF(COUNTIFS($A$2:A{r},A{r},$B$2:B{r},B{r})=1,1,0))')
    S.cell(r, 7, f'=IF(AND(A{r}<>"",D{r}="Active",F{r}=1),COUNTIFS($B$2:$B${send},B{r},$D$2:$D${send},"Active",$F$2:$F${send},1,$A$2:$A${send},"<"&A{r})+1,"")')
    S.cell(r, 8, f'=IF(G{r}="","",B{r}&"|"&G{r})')
    for c in range(1, 5): S.cell(r, c).font = BF
for c in 'FGH': S.column_dimensions[c].hidden = True
for c, w in zip('ABCD', (24, 28, 14, 11)): S.column_dimensions[c].width = w
S['I1'] = 'How to use'; S['I1'].font = Font(name='Arial', bold=True)
for i, t in enumerate([
    'One row per staff member per university (and campus where it matters).',
    'Only Active staff appear in the Staff Member dropdown on that university’s visit table.',
    'When someone leaves, set Status to Inactive. Don’t delete the row; that keeps history intact.',
    f'Spare rows are ready to the bottom of the table (row {send}).'], 2):
    S.cell(i, 9, t).font = NOTE
dv(S, [f'D2:D{send}'], type='list', formula1='"Active,Inactive"')
dv(S, [f'B2:B{send}'], type='list', formula1='"' + ','.join(u[1] for u in UNIS) + '"')
S.freeze_panes = 'A2'
SM = "'Staff Management'"

# ------------------------------------------------------------------ University tabs
uni_info = {}
VROWS = 500
for tab, full, pre, tstem, nstem, mlabel in UNIS:
    ws = wb.create_sheet(tab)
    ws.sheet_properties.tabColor = NAVY
    accs = accounts.get(mlabel, []) if mlabel else []
    # unique names
    seen = {}
    for a in accs:
        k = a['name'].lower(); seen[k] = seen.get(k, 0) + 1
        if seen[k] > 1: a['name'] = f"{a['name']} ({seen[k]})"
    nloc = len(accs); LSPARE = 30
    L0 = 5; lfirst = L0 + 1; llast = L0 + max(nloc, 0) + LSPARE
    V0 = llast + 4; vfirst = V0 + 1; vlast = V0 + VROWS

    ws['A1'] = full; ws['A1'].font = Font(name='Arial', bold=True, size=16)
    ws['A2'] = ('Locations: every store we visit. Add new ones in the blank rows (give each a new ID, never reuse one). '
                'Set Status to Inactive to stop visiting a location without losing its history.')
    ws['A3'] = ('Visits: log each visit further down. Only Completed visits count and reset the 14-day clock. '
                'Grey columns calculate themselves.')
    ws['A2'].font = NOTE; ws['A3'].font = NOTE
    ws.cell(L0 - 1, 1, 'LOCATIONS  (fortnightly reminder + progress)').font = Font(name='Arial', bold=True, size=11, color=NAVY)

    lcols = ['Location ID', 'Store/Location', 'Campus', 'Address', 'Status', 'Date Added', 'Last Visit Date',
             'Next Visit Due', 'Days Until Due', 'Visit Status', 'Completed Visits', 'DPs Added', 'DPs Lost',
             'Net DPs', 'Editions Sold In', 'POS Placed', 'Coolers Installed', 'Coolers Removed']
    for i, h in enumerate(lcols, 1): ws.cell(L0, i, h)
    vR = lambda c: f'${c}${vfirst}:${c}${vlast}'
    for j in range(llast - L0):
        r = lfirst + j
        if j < nloc:
            a = accs[j]
            ws.cell(r, 1, f'{pre}-{j + 1:03d}'); ws.cell(r, 2, a['name']); ws.cell(r, 3, a['campus'] or None)
            ws.cell(r, 4, a['addr'] or None); ws.cell(r, 5, 'Active'); ws.cell(r, 6, TODAY)
            a['id'] = f'{pre}-{j + 1:03d}'
        A = f'$A{r}'
        ws.cell(r, 7, f'=IF({A}="","",IF(COUNTIFS({vR("B")},{A},{vR("F")},"Completed")=0,"",_xlfn.MAXIFS({vR("D")},{vR("B")},{A},{vR("F")},"Completed")))')
        ws.cell(r, 8, f'=IF({A}="","",IF(G{r}="",IF(F{r}="","",F{r}+14),G{r}+14))')
        ws.cell(r, 9, f'=IF(OR({A}="",H{r}=""),"",H{r}-TODAY())')
        ws.cell(r, 10, f'=IF({A}="","",IF(E{r}="Inactive","Inactive",IF(G{r}="","No Visit Recorded",IF(H{r}<TODAY(),"Overdue",IF(H{r}<=TODAY()+3,"Due Soon","On Track")))))')
        ws.cell(r, 11, f'=IF({A}="","",COUNTIFS({vR("B")},{A},{vR("F")},"Completed"))')
        for c, vc in [(12, 'H'), (13, 'I'), (15, 'K'), (16, 'M'), (17, 'O'), (18, 'P')]:
            ws.cell(r, c, f'=IF({A}="","",SUMIFS({vR(vc)},{vR("B")},{A},{vR("F")},"Completed"))')
        ws.cell(r, 14, f'=IF({A}="","",L{r}-M{r})')
        # helper rank for active location list
        ws.cell(r, 29, f'=IF(AND(E{r}="Active",B{r}<>""),COUNTIFS($E${lfirst}:$E${llast},"Active",$B${lfirst}:$B${llast},"<"&B{r})+1,"")')
        for c in range(1, 19):
            x = ws.cell(r, c); x.font = BF; x.border = B
            if c >= 7: x.fill = AUTO; x.alignment = Alignment(horizontal='center')
        for c in (6, 7, 8): ws.cell(r, c).number_format = 'dd/mm/yyyy'
        ws.cell(r, 14).number_format = '+0;-0;0'
    table(ws, f'tbl{tstem}Locations', f'A{L0}:R{llast}')

    ws.cell(V0 - 1, 1, 'VISITS  (one row per visit)').font = Font(name='Arial', bold=True, size=11, color=NAVY)
    vcols = ['Visit ID', 'Location ID', 'Store/Location', 'Visit Date', 'Staff Member', 'Visit Status', 'Notes',
             'DPs Added', 'DPs Lost', 'Net DPs', 'Editions Sold In', 'Editions Flavours', 'POS Placed', 'POS Type',
             'Coolers Installed', 'Coolers Removed', 'Next Step']
    for i, h in enumerate(vcols, 1): ws.cell(V0, i, h)
    seeds = [a for a in accs if a['last']]
    seeds.sort(key=lambda a: a['last'])
    lR = lambda c: f'${c}${lfirst}:${c}${llast}'
    for j in range(VROWS):
        r = vfirst + j
        if j < len(seeds):
            a = seeds[j]
            ws.cell(r, 3, a['name']); ws.cell(r, 4, a['last']); ws.cell(r, 6, 'Completed')
            ws.cell(r, 7, 'Imported: last visit recorded on the audit tab')
            if a['nxt']: ws.cell(r, 17, a['nxt'])
        ws.cell(r, 1, f'=IF(C{r}="","","{pre}-V"&TEXT(ROW()-{V0},"0000"))')
        ws.cell(r, 2, f'=IF(C{r}="","",IFERROR(INDEX({lR("A")},MATCH(C{r},{lR("B")},0)),"Not found"))')
        ws.cell(r, 10, f'=IF(AND(H{r}="",I{r}=""),"",N(H{r})-N(I{r}))')
        # helper: first completed visit for this location in this month (for distinct counts)
        ws.cell(r, 30, f'=IF(OR(B{r}="",D{r}="",F{r}<>"Completed"),"",IF(COUNTIFS($B${vfirst}:B{r},B{r},$F${vfirst}:F{r},"Completed",$D${vfirst}:D{r},">="&(EOMONTH(D{r},-1)+1),$D${vfirst}:D{r},"<="&EOMONTH(D{r},0))=1,1,0))')
        for c in range(1, 18):
            x = ws.cell(r, c); x.font = BF; x.border = B
            if c in (1, 2, 10): x.fill = AUTO
        ws.cell(r, 4).number_format = 'dd/mm/yyyy'; ws.cell(r, 10).number_format = '+0;-0;0'
    table(ws, f'tbl{tstem}Visits', f'A{V0}:Q{vlast}')

    # helper lists (hidden AA:AD)
    ws['AA1'] = 'Helper: active staff'; ws['AB1'] = 'Helper: active locations'; ws['AC1'] = 'Helper: location rank'
    ws['AD1'] = 'Helper: first visit of month'
    for k in range(1, 41):
        r = k + 1
        ws.cell(r, 27, f'=IFERROR(INDEX({SM}!$A$2:$A${send},MATCH("{full}|{k}",{SM}!$H$2:$H${send},0)),"")')
    nslots = llast - L0
    for k in range(1, nslots + 1):
        ws.cell(k + 1, 28, f'=IFERROR(INDEX({lR("B")},MATCH({k},{lR("AC")},0)),"")')
    for c in ('AA', 'AB', 'AC', 'AD'): ws.column_dimensions[c].hidden = True

    staff_name = f'{nstem}StaffList'; loc_name = f'{nstem}LocationList'
    wb.defined_names[staff_name] = DefinedName(staff_name, attr_text=f'OFFSET({q(tab)}!$AA$2,0,0,MAX(1,COUNTIF({q(tab)}!$AA$2:$AA$41,"?*")),1)')
    wb.defined_names[loc_name] = DefinedName(loc_name, attr_text=f'OFFSET({q(tab)}!$AB$2,0,0,MAX(1,COUNTIF({q(tab)}!$AB$2:$AB${nslots + 1},"?*")),1)')

    dv(ws, [f'E{lfirst}:E{llast}'], type='list', formula1='"Active,Inactive"')
    dv(ws, [f'F{lfirst}:F{llast}', f'D{vfirst}:D{vlast}'], type='date', operator='greaterThan', formula1='DATE(2020,1,1)', error='Enter a date, e.g. 08/10/2026')
    dv(ws, [f'C{vfirst}:C{vlast}'], type='list', formula1=f'={loc_name}', error='Pick an active location from the list.')
    dv(ws, [f'E{vfirst}:E{vlast}'], type='list', formula1=f'={staff_name}', error='Pick a staff member allocated to this university (see Staff Management).')
    dv(ws, [f'F{vfirst}:F{vlast}'], type='list', formula1='"Scheduled,Completed,Cancelled"')
    dv(ws, [f'N{vfirst}:N{vlast}'], type='list', formula1='"Price tag,Wobbler,Poster,Shelf strip,Fridge decal,Counter unit,Menu board,Other"')
    dv(ws, [f'{c}{vfirst}:{c}{vlast}' for c in 'HIKMOP'], type='whole', operator='between', formula1='0', formula2='200', error='Enter a whole number (0 or more).')

    st = f'J{lfirst}:J{llast}'
    ws.conditional_formatting.add(st, FormulaRule(formula=[f'$J{lfirst}="Overdue"'], fill=PatternFill('solid', fgColor='9C0006'), font=Font(color='FFFFFF', bold=True)))
    ws.conditional_formatting.add(st, FormulaRule(formula=[f'$J{lfirst}="Due Soon"'], fill=PatternFill('solid', fgColor='FFC000'), font=Font(color='000000')))
    ws.conditional_formatting.add(st, FormulaRule(formula=[f'$J{lfirst}="No Visit Recorded"'], fill=PatternFill('solid', fgColor='FFF2CC'), font=Font(bold=True)))
    ws.conditional_formatting.add(st, FormulaRule(formula=[f'$J{lfirst}="On Track"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    ws.conditional_formatting.add(f'N{lfirst}:N{llast}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color='C00000', bold=True)))
    ws.conditional_formatting.add(f'J{vfirst}:J{vlast}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color='C00000', bold=True)))

    widths = [12, 30, 13, 30, 10, 12, 13, 13, 10, 18, 11, 9, 9, 9, 11, 10, 11, 11]
    for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
    ws.column_dimensions['G'].width = 34  # also Notes in the visit table
    ws.column_dimensions['Q'].width = 40  # Next Step
    ws.row_dimensions[L0].height = 30; ws.row_dimensions[V0].height = 30
    ws.freeze_panes = 'C6'
    uni_info[tab] = dict(full=full, pre=pre, lfirst=lfirst, llast=llast, vfirst=vfirst, vlast=vlast, mlabel=mlabel, n=nloc, seeds=len(seeds))

# ------------------------------------------------------------------ VIC Master: progress block from university tabs
cNS = mh['Next Step']; P0 = cNS + 1
lab2tab = {u[5]: u[0] for u in UNIS if u[5]}
pcols = [('Location ID', None), ('Visit Status', 'J'), ('Last Visit Date', 'G'), ('Next Visit Due', 'H'), ('Days Until Due', 'I'),
         ('Completed Visits', 'K'), ('DPs Added', 'L'), ('DPs Lost', 'M'), ('Net DPs', 'N'), ('Editions Sold In', 'O'),
         ('POS Placed', 'P'), ('Coolers Installed', 'Q'), ('Coolers Removed', 'R')]
for c in range(P0, P0 + 20):
    for r in range(MH - 1, mlast + 1): M.cell(r, c).value = None
for rng in list(M.merged_cells.ranges):
    if rng.min_row == MH - 1: M.unmerge_cells(str(rng))
M.merge_cells(start_row=MH - 1, start_column=4, end_row=MH - 1, end_column=cNS)
head(M, MH - 1, 4, 'ACCOUNT DETAILS  (live from each audit tab)')
M.merge_cells(start_row=MH - 1, start_column=P0, end_row=MH - 1, end_column=P0 + len(pcols) - 1)
head(M, MH - 1, P0, 'ACCOUNT PROGRESS  (live from each university tab’s Locations table)', GREEN)
for i, (h, _) in enumerate(pcols): head(M, MH, P0 + i, h, GREEN)
idx_by_row = {a['row']: a.get('id') for lst in accounts.values() for a in lst}
for r in range(mfirst, mlast + 1):
    tab = lab2tab[M.cell(r, 2).value]; ui = uni_info[tab]; t = q(tab)
    M.cell(r, P0, idx_by_row[r])
    for i, (h, col) in enumerate(pcols[1:], 1):
        M.cell(r, P0 + i, f'=IFERROR(INDEX({t}!${col}${ui["lfirst"]}:${col}${ui["llast"]},MATCH(${L(P0)}{r},{t}!$A${ui["lfirst"]}:$A${ui["llast"]},0)),"")')
    for i in range(len(pcols)):
        x = M.cell(r, P0 + i); x.font = BF; x.border = B
        if i: x.fill = PatternFill('solid', fgColor='E2EFDA'); x.alignment = Alignment(horizontal='center')
    for i in (2, 3): M.cell(r, P0 + i).number_format = 'dd/mm/yyyy'
    M.cell(r, P0 + 8).number_format = '+0;-0;0'
for i, w in enumerate([11, 18, 12, 12, 10, 10, 9, 9, 9, 10, 9, 10, 10]): M.column_dimensions[L(P0 + i)].width = w
vs = f'{L(P0 + 1)}{mfirst}:{L(P0 + 1)}{mlast}'; c0 = f'${L(P0 + 1)}{mfirst}'
M.conditional_formatting.add(vs, FormulaRule(formula=[f'{c0}="Overdue"'], fill=PatternFill('solid', fgColor='9C0006'), font=Font(color='FFFFFF', bold=True)))
M.conditional_formatting.add(vs, FormulaRule(formula=[f'{c0}="Due Soon"'], fill=PatternFill('solid', fgColor='FFC000')))
M.conditional_formatting.add(vs, FormulaRule(formula=[f'{c0}="No Visit Recorded"'], fill=PatternFill('solid', fgColor='FFF2CC'), font=Font(bold=True)))
M.conditional_formatting.add(vs, FormulaRule(formula=[f'{c0}="On Track"'], fill=PatternFill('solid', fgColor='C6EFCE')))
M.auto_filter.ref = f'A{MH}:{L(P0 + len(pcols) - 1)}{mlast}'
M.freeze_panes = M.cell(MH + 1, mh['Account'] + 1)
M['A3'] = 'Green Account Progress columns come from each university tab (visit status uses the 14-day cycle). Log visits on the university tabs, not here.'
M['A3'].font = Font(name='Arial', italic=True, size=9, color=GREEN)

# ------------------------------------------------------------------ Progress tab
G = wb.create_sheet('Progress'); G.sheet_properties.tabColor = GREEN
G['A1'] = 'University Visit Progress'; G['A1'].font = Font(name='Arial', bold=True, size=16)
G['A2'] = '=" Fortnightly snapshot as at "&TEXT(TODAY(),"dd mmm yyyy")&". A location counts as completed if it had a Completed visit in the last 14 days."'
G['A2'].font = NOTE
pc = ['University', 'Active Locations', 'Required Visits', 'Completed Visits', 'Completion %', 'On Track', 'Overdue', 'No Visit Recorded']
P_H = 4
for i, h in enumerate(pc, 1): G.cell(P_H, i, h)
for j, (tab, full, *_r) in enumerate(UNIS):
    r = P_H + 1 + j; ui = uni_info[tab]; t = q(tab)
    E = f'{t}!$E${ui["lfirst"]}:$E${ui["llast"]}'; Gc = f'{t}!$G${ui["lfirst"]}:$G${ui["llast"]}'; J = f'{t}!$J${ui["lfirst"]}:$J${ui["llast"]}'
    G.cell(r, 1, full)
    G.cell(r, 2, f'=COUNTIF({E},"Active")')
    G.cell(r, 3, f'=B{r}')
    G.cell(r, 4, f'=COUNTIFS({E},"Active",{Gc},">="&TODAY()-14,{Gc},"<="&TODAY())')
    G.cell(r, 5, f'=IFERROR(D{r}/C{r},0)').number_format = '0%'
    G.cell(r, 6, f'=COUNTIFS({E},"Active",{J},"On Track")+COUNTIFS({E},"Active",{J},"Due Soon")')
    G.cell(r, 7, f'=COUNTIFS({E},"Active",{J},"Overdue")')
    G.cell(r, 8, f'=COUNTIFS({E},"Active",{J},"No Visit Recorded")')
    for c in range(1, 9):
        x = G.cell(r, c); x.font = BF
        if c > 1: x.alignment = Alignment(horizontal='center')
P_L = P_H + len(UNIS)
table(G, 'tblProgress', f'A{P_H}:H{P_L}', 'TableStyleMedium2')
tr = P_L + 1
G.cell(tr, 1, 'Total').font = Font(name='Arial', bold=True)
for c in range(2, 9):
    if c == 5: G.cell(tr, c, f'=IFERROR(D{tr}/C{tr},0)').number_format = '0%'
    else: G.cell(tr, c, f'=SUM({L(c)}{P_H + 1}:{L(c)}{P_L})')
    G.cell(tr, c).font = Font(name='Arial', bold=True); G.cell(tr, c).alignment = Alignment(horizontal='center')
for c in range(1, 9): G.cell(tr, c).fill = PatternFill('solid', fgColor='D9D9D9')
G.cell(tr + 1, 1, 'On Track includes Due Soon (within 3 days). No Visit Recorded is shown separately; add it to Overdue for a stricter view.').font = NOTE
G.conditional_formatting.add(f'E{P_H + 1}:E{tr}', FormulaRule(formula=[f'AND(C{P_H + 1}>0,E{P_H + 1}<0.5)'], font=Font(color='C00000', bold=True)))

# outcomes by university
O_H = tr + 4
G.cell(O_H - 1, 1, 'Outcomes by university (all Completed visits)').font = Font(name='Arial', bold=True, size=11, color=NAVY)
oc = [('University', None), ('Completed Visits', 'K'), ('DPs Added', 'L'), ('DPs Lost', 'M'), ('Net DPs', None), ('Editions Sold In', 'O'),
      ('POS Placed', 'P'), ('Coolers Installed', 'Q'), ('Coolers Removed', 'R'), ('Net Coolers', None)]
for i, (h, _) in enumerate(oc, 1): head(G, O_H, i, h)
for j, (tab, full, *_r) in enumerate(UNIS):
    r = O_H + 1 + j; ui = uni_info[tab]; t = q(tab)
    G.cell(r, 1, full)
    for i, (h, col) in enumerate(oc[1:], 2):
        if col: G.cell(r, i, f'=SUM({t}!${col}${ui["lfirst"]}:${col}${ui["llast"]})')
    G.cell(r, 5, f'=C{r}-D{r}'); G.cell(r, 10, f'=H{r}-I{r}')
ot = O_H + len(UNIS) + 1
G.cell(ot, 1, 'Total')
for c in range(2, 11): G.cell(ot, c, f'=SUM({L(c)}{O_H + 1}:{L(c)}{ot - 1})')
for r in range(O_H + 1, ot + 1):
    for c in range(1, 11):
        x = G.cell(r, c); x.border = B; x.font = Font(name='Arial', size=10, bold=(r == ot))
        if c > 1: x.alignment = Alignment(horizontal='center')
        if r == ot: x.fill = PatternFill('solid', fgColor='D9D9D9')
    G.cell(r, 5).number_format = G.cell(r, 10).number_format = '+0;-0;0'

# by month
Y_R = ot + 3
G.cell(Y_R, 1, 'Outcomes by month').font = Font(name='Arial', bold=True, size=11, color=NAVY)
G.cell(Y_R, 3, 'Year'); G.cell(Y_R, 3).font = Font(name='Arial', bold=True)
G.cell(Y_R, 4, 2026); G.cell(Y_R, 4).font = Font(name='Arial', bold=True, color='0000FF'); G.cell(Y_R, 4).fill = PatternFill('solid', fgColor='FFF2CC'); G.cell(Y_R, 4).border = B
G.cell(Y_R, 5, '← change the year').font = NOTE
mc = [('Month', None), ('Completed Visits', None), ('Locations Visited', 'AD'), ('DPs Added', 'H'), ('DPs Lost', 'I'), ('Net DPs', None),
      ('Editions Sold In', 'K'), ('POS Placed', 'M'), ('Coolers Installed', 'O'), ('Coolers Removed', 'P'), ('Net Coolers', None)]
MH_ = Y_R + 1
for i, (h, _) in enumerate(mc, 1): head(G, MH_, i, h)
for m in range(1, 13):
    r = MH_ + m
    G.cell(r, 1, f'=DATE(${L(4)}${Y_R},{m},1)').number_format = 'mmm yyyy'
    def summ(col):
        parts = []
        for tab, *_r in UNIS:
            ui = uni_info[tab]; t = q(tab)
            D = f'{t}!$D${ui["vfirst"]}:$D${ui["vlast"]}'; F = f'{t}!$F${ui["vfirst"]}:$F${ui["vlast"]}'
            crit = f'{D},">="&$A{r},{D},"<"&EDATE($A{r},1)'
            if col is None: parts.append(f'COUNTIFS({F},"Completed",{crit})')
            else: parts.append(f'SUMIFS({t}!${col}${ui["vfirst"]}:${col}${ui["vlast"]},{F},"Completed",{crit})')
        return '=' + '+'.join(parts)
    G.cell(r, 2, summ(None))
    for i, (h, col) in enumerate(mc[2:], 3):
        if col: G.cell(r, i, summ(col))
    G.cell(r, 6, f'=D{r}-E{r}'); G.cell(r, 11, f'=I{r}-J{r}')
mt = MH_ + 13
G.cell(mt, 1, 'Year total')
for c in range(2, 12): G.cell(mt, c, f'=SUM({L(c)}{MH_ + 1}:{L(c)}{mt - 1})')
for r in range(MH_ + 1, mt + 1):
    for c in range(1, 12):
        x = G.cell(r, c); x.border = B; x.font = Font(name='Arial', size=10, bold=(r == mt))
        if c > 1: x.alignment = Alignment(horizontal='center')
        if r == mt: x.fill = PatternFill('solid', fgColor='D9D9D9')
    G.cell(r, 6).number_format = G.cell(r, 11).number_format = '+0;-0;0'
G.cell(mt + 1, 1, 'Locations Visited = distinct locations with a Completed visit that month. The year total adds the months, so a location visited in two months counts twice.').font = NOTE
G.column_dimensions['A'].width = 26
for c in range(2, 12): G.column_dimensions[L(c)].width = 12
G.row_dimensions[P_H].height = 30; G.row_dimensions[O_H].height = 30; G.row_dimensions[MH_].height = 30
for rr in (P_H,):
    for c in range(1, 9): G.cell(rr, c).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def bar(title, cols, colors, anchor):
    ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'; ch.title = title
    ch.y_axis.title = 'Number of Locations'; ch.x_axis.title = None
    data = Reference(G, min_col=cols[0], max_col=cols[-1], min_row=P_H, max_row=P_L) if cols == list(range(cols[0], cols[-1] + 1)) else None
    for col, color in zip(cols, colors):
        ch.add_data(Reference(G, min_col=col, min_row=P_H, max_row=P_L), titles_from_data=True)
        s = ch.series[-1]; s.graphicalProperties.solidFill = color; s.graphicalProperties.line.solidFill = color
    ch.set_categories(Reference(G, min_col=1, min_row=P_H + 1, max_row=P_L))
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.x_axis.scaling.orientation = 'maxMin'
    ch.y_axis.majorGridlines = None; ch.legend.position = 'b'
    ch.width = 17; ch.height = 9.5; ch.gapWidth = 60
    G.add_chart(ch, anchor)
bar('Fortnightly University Visit Progress', [3, 4], ['BFBFBF', '2F5597'], 'M3')
bar('Location Visit Status by University', [6, 7], ['70AD47', 'C00000'], 'M23')

# ------------------------------------------------------------------ order + Index
order = ['Index', 'Progress', 'VIC Master', 'VIC Master (Vending)', 'Staff Management']
groups = [('Monash', ['Monash Clayton Audit', 'Monash Caulfield Audit']), ('La Trobe', ['La Trobe Audit', 'La Trobe Vending']),
          ('Deakin', ['Deakin Audit', 'Deakin Vending']), ('Victoria University', ['Victoria Uni Audit', 'Victoria Uni Vending']),
          ('RMIT', ['RMIT Audit', 'RMIT Vending']), ('Swinburne', ['Swinburne Audit', 'Swinburne Vending']),
          ('Melbourne University', ['Melbourne Uni Audit', 'Melbourne Uni Vending']), ('UTAS', [])]
for t, extra in groups: order += [t] + extra
order += ['Template Audit', 'Template Vending', 'Calendar']
assert sorted(order) == sorted(wb.sheetnames), (set(order) ^ set(wb.sheetnames))
wb._sheets = [wb[n] for n in order]
for n in [g for g, _ in groups[:-1]]:
    pass
for t, extra in groups:
    for e in extra: wb[e].sheet_properties.tabColor = '8EA9DB'

I = wb['Index']
for row in I.iter_rows(min_row=1):
    for c in row: c.value = None; c.hyperlink = None
I['A1'] = 'Education Channel Execution Tracker'; I['A1'].font = Font(name='Arial', bold=True, size=14)
I['A2'] = 'Click a tab name to jump to it. Log visits on each university tab; everything else updates from there.'; I['A2'].font = NOTE
for i, h in enumerate(['Group', 'Tab', 'What it’s for'], 1): head(I, 4, i, h)
desc = {'Progress': 'Fortnightly visit progress, outcomes by university and month, charts',
        'VIC Master': 'All VIC on-premise accounts, audit data + live progress',
        'VIC Master (Vending)': 'All VIC vending machines', 'Staff Management': 'Who can be picked for visits at each university',
        'Calendar': 'Month calendar (from the Uni Melb and RMIT files)', 'Template Audit': 'Blank audit template', 'Template Vending': 'Blank vending template'}
r = 5
for n in order[1:]:
    grp = 'Overview' if n in ('Progress', 'VIC Master', 'VIC Master (Vending)', 'Staff Management') else ''
    d = desc.get(n)
    if n in uni_info:
        ui = uni_info[n]; grp = ui['full']
        d = f"Locations ({ui['n']}) with 14-day reminder + visit log"
    elif d is None:
        d = 'Ranging audit (from original workbook)' if 'Audit' in n else 'Vending audit (from original workbook)'
    if n.startswith('Template'): grp = 'Template'
    I.cell(r, 1, grp).font = Font(name='Arial', bold=bool(grp))
    x = I.cell(r, 2, n); x.hyperlink = f"#{q(n)}!A1"; x.font = Font(name='Arial', color='0563C1', underline='single', bold=(n in uni_info or grp == 'Overview'))
    I.cell(r, 3, d).font = Font(name='Arial', color='595959')
    r += 1
I.column_dimensions['A'].width = 24; I.column_dimensions['B'].width = 26; I.column_dimensions['C'].width = 62
wb.active = 0
for w in wb.worksheets: w.sheet_view.tabSelected = (w.title == 'Index')
wb.save(p)
for k, v in uni_info.items(): print(k, v)
