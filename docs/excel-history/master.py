import sys, openpyxl, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
p=sys.argv[1]; wb=openpyxl.load_workbook(p)
for n in ['VIC Master','VIC Master (Vending)']:
    if n in wb.sheetnames: del wb[n]
norm=lambda s: re.sub(r'[^a-z0-9#]','',str(s).lower()) if s is not None else ''
ONP=[('Deakin','Deakin Audit',{3}),('La Trobe','La Trobe Audit',{3}),('Melbourne University','Melbourne Uni Audit',set()),('Monash','Monash Clayton Audit',set()),('Monash','Monash Caulfield Audit',set()),('RMIT','RMIT Audit',set()),('Swinburne','Swinburne Audit',set()),('Victoria University','Victoria Uni Audit',set())]
VEND=[('Deakin','Deakin Vending',{3}),('La Trobe','La Trobe Vending',{3}),('Melbourne University','Melbourne Uni Vending',set()),('RMIT','RMIT Vending',{3}),('Swinburne','Swinburne Vending',set()),('Victoria University','Victoria Uni Vending',set())]
UNIS=['Monash','La Trobe','Deakin','Victoria University','RMIT','Swinburne','Melbourne University']
ONP_COLS=[('campus','Campus'),('bull',"Bullring / Bull's Eye"),('account','Account'),('city','City'),('street','Street'),('landmark','Landmark'),('postcode','Postcode'),
 ('ed','ED'),('zero','Zero'),('sf','SF'),('edition','Edition'),('monster','Monster'),('monsterultra','Monster Ultra'),('v','V'),('vsf','V SF'),('vblue','V Blue'),('soloenergy','Solo Energy'),
 ('pos','POS'),('cooler','Cooler'),('channel','Channel'),('classification','Classification'),('indnka','IND / NKA'),('pic','Person-in-charge'),('contact','Contact'),('price','Price of Can'),('topcomp','Top Competitor'),('comppricing','Competitor Pricing'),('lastvisit','Last Visit'),('nextstep','Next Step')]
ALIAS={'campus':'campus','bullringbullseye':'bull','account':'account','city':'city','street':'street','landmark':'landmark','postcode':'postcode','pos':'pos','cooler':'cooler','channel':'channel','classification':'classification','indnka':'indnka','personincharge':'pic','personic':'pic','contact':'contact','priceofcan':'price','topcompetitor':'topcomp','competitorpricing':'comppricing','lastvisit':'lastvisit','nextstep':'nextstep',
 'ed':'ed','zero':'zero','sf':'sf','edition':'edition','editions':'edition','monster':'monster','monsterultra':'monsterultra','v':'v','vsf':'vsf','vblue':'vblue','soloenergy':'soloenergy'}
def onp_map(ws):
    m={}
    for c in range(1,ws.max_column+1):
        for r in (1,2):
            k=ALIAS.get(norm(ws.cell(r,c).value))
            if k and k not in m and not (r==1 and k in('ed','sf','v')): m[k]=c
    return m
VEND_COLS=[('campus','Campus'),('building','Building'),('level','Level'),('landmark','Landmark'),('location','Location'),('company','Vending Company'),('machine','Machine ID'),
 ('ed','ED'),('sf','SF'),('editions','Editions'),('rbprice','Red Bull Price'),('vr','V Ranging'),('vsku','V # of SKU'),('vprice','V Price'),('mr','Monster Ranging'),('msku','Monster # of SKU'),('mprice','Monster Price')]
def vend_map(ws):
    m={}; prices=[]; rng=[]; sku=[]
    for c in range(1,ws.max_column+1):
        for r in (1,2):
            n=norm(ws.cell(r,c).value)
            if n in('campus','rmitcampus'): m['campus']=c
            elif n in('building','level','landmark','location'): m.setdefault(n,c)
            elif n=='vendingcompany': m['company']=c
            elif n=='machineid': m['machine']=c
            elif r==2 and n in('ed','sf'): m[n]=c
            elif r==2 and n=='editions': m['editions']=c
            elif r==2 and n=='price': prices.append(c)
            elif r==2 and n=='ranging': rng.append(c)
            elif r==2 and n=='#ofsku': sku.append(c)
    for k,v in zip(['rbprice','vprice','mprice'],prices): m[k]=v
    for k,v in zip(['vr','mr'],rng): m[k]=v
    for k,v in zip(['vsku','msku'],sku): m[k]=v
    return m
def data_rows(ws,keycols,skip):
    out=[]
    for r in range(3,ws.max_row+1):
        vals=[ws.cell(r,c).value for c in range(1,ws.max_column+1)]
        if any(isinstance(x,str) and x.startswith('=COUNTIF') for x in vals): continue
        if norm(vals[0]) in('sm','mts','both','sample') and all(v is None for v in vals[1:]): continue
        if r in skip: continue
        if any(ws.cell(r,c).value not in(None,'',' ','\xa0') for c in keycols): out.append(r)
    return out
NAVY='1F3864'; thin=Side(style='thin',color='BFBFBF'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
HF=Font(name='Arial',bold=True,color='FFFFFF',size=10); BF=Font(name='Arial',size=10)
GRP={'ed':'D9E1F2','zero':'D9E1F2','sf':'D9E1F2','edition':'D9E1F2','editions':'D9E1F2','rbprice':'D9E1F2','monster':'FCE4D6','monsterultra':'FCE4D6','v':'FCE4D6','vsf':'FCE4D6','vblue':'FCE4D6','soloenergy':'FCE4D6','vr':'FCE4D6','vsku':'FCE4D6','vprice':'FCE4D6','mr':'FCE4D6','msku':'FCE4D6','mprice':'FCE4D6'}
def build(name,specs,cols,mapper,keyfields,start_row,pct_fields,title,note):
    ws=wb.create_sheet(name,1 if 'Vending' not in name else 2)
    ws['A1']=title; ws['A1'].font=Font(name='Arial',bold=True,size=14)
    ws['A2']=note; ws['A2'].font=Font(name='Arial',italic=True,size=9,color='666666')
    hdr=['State','University','Source tab']+[h for _,h in cols]
    H=start_row
    for i,h in enumerate(hdr,1):
        c=ws.cell(H,i,h); c.font=HF; c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=B
    for i,(k,_) in enumerate(cols,4):
        if k in GRP: ws.cell(H,i).fill=PatternFill('solid',fgColor='2F5597' if GRP[k]=='D9E1F2' else 'C55A11')
    r=H+1; counts={}
    for uni,tab,skip in specs:
        src=wb[tab]; m=mapper(src); keys=[m[k] for k in keyfields if k in m]
        rows=data_rows(src,keys,skip); counts[tab]=len(rows)
        for sr in rows:
            ws.cell(r,1,'VIC'); ws.cell(r,2,uni); ws.cell(r,3,tab)
            for i,(k,_) in enumerate(cols,4):
                if k in m:
                    ref=f"'{tab}'!{L(m[k])}{sr}"
                    ws.cell(r,i,f'=IF({ref}="","",{ref})')
                    if k=='lastvisit': ws.cell(r,i).number_format='dd/mm/yyyy'
                    if k in('price','comppricing','rbprice','vprice','mprice'): ws.cell(r,i).number_format='$0.00'
            for i in range(1,len(hdr)+1):
                c=ws.cell(r,i); c.font=BF; c.border=B
                k=cols[i-4][0] if i>=4 else None
                if k in GRP: c.fill=PatternFill('solid',fgColor=GRP[k]); c.alignment=Alignment(horizontal='center')
            r+=1
    last=r-1
    ws.auto_filter.ref=f"A{H}:{L(len(hdr))}{last}"
    ws.freeze_panes=ws.cell(H+1,6)
    ws.row_dimensions[H].height=30
    for i,h in enumerate(hdr,1):
        w=max(9,min(40,len(h)+3))
        if h in('Account','Street','Landmark','Machine ID','Location'): w=26
        if h=='Next Step': w=60
        if h in('University','Source tab','Campus','Vending Company'): w=max(w,16)
        if h=='Source tab': w=26
        ws.column_dimensions[L(i)].width=w
    # summary
    colof={k:L(i) for i,(k,_) in enumerate(cols,4)}
    U=f"$B${H+1}:$B${last}"; ACC=f"${colof[keyfields[0]]}${H+1}:${colof[keyfields[0]]}${last}"
    sr=4; ws.cell(sr,1,'Summary by university').font=Font(name='Arial',bold=True,size=11)
    shd=['University',' # '+('accounts' if 'Vending' not in name else 'machines')]+[lab for _,lab in pct_fields]
    for i,h in enumerate(shd,2):
        c=ws.cell(sr+1,i,h); c.font=HF; c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(horizontal='center',wrap_text=True); c.border=B
    ws.row_dimensions[sr+1].height=30
    for j,u in enumerate(UNIS+['VIC total']):
        rr=sr+2+j; ws.cell(rr,2,u)
        crit= f'({U}="{u}")*' if u!='VIC total' else ''
        ws.cell(rr,3,(f'=COUNTIF({U},"{u}")' if u!='VIC total' else f'=SUM(C{sr+2}:C{sr+1+len(UNIS)})'))
        for i,(k,_) in enumerate(pct_fields,4):
            col=f"${colof[k]}${H+1}:${colof[k]}${last}"
            if k in('price','rbprice'):
                f=f'=IFERROR(AVERAGEIFS({col},{U},"{u}",{col},">0"),"-")' if u!='VIC total' else f'=IFERROR(AVERAGEIF({col},">0"),"-")'
                ws.cell(rr,i,f).number_format='$0.00'
            else:
                ws.cell(rr,i,f'=IF(C{rr}=0,"-",SUMPRODUCT({crit}({col}=1))/C{rr})').number_format='0%'
        for i in range(2,len(shd)+2):
            c=ws.cell(rr,i); c.border=B; c.font=Font(name='Arial',size=10,bold=(u=='VIC total'))
            if i>2: c.alignment=Alignment(horizontal='center')
            if u=='VIC total': c.fill=PatternFill('solid',fgColor='D9D9D9')
    ws.cell(sr+2+len(UNIS)+1,2,'% = share of rows marked 1; blanks count as not ranged. Avg prices skip $0 (not stocked) and text entries.').font=Font(name='Arial',italic=True,size=9,color='666666')
    return counts,last-H
c1,n1=build('VIC Master',ONP,ONP_COLS,onp_map,['account'],17,
  [('ed','ED ranged'),('sf','SF ranged'),('edition','Edition ranged'),('pos','POS'),('cooler','Cooler'),('monster','Monster ranged'),('v','V ranged'),('soloenergy','Solo ranged'),('price','Avg can price')],
  'VIC Master: on-premise accounts, all universities','Live: every cell links back to its university tab, so edits there show here. A new account added to a university tab needs a row added here too. Template example rows are excluded.')
c2,n2=build('VIC Master (Vending)',VEND,VEND_COLS,vend_map,['machine','building','company'],17,
  [('ed','ED ranged'),('sf','SF ranged'),('editions','Editions ranged'),('vr','V ranged'),('mr','Monster ranged'),('rbprice','Avg Red Bull price')],
  'VIC Master (Vending): machines, all universities','Live: every cell links back to its university vending tab. A new machine added to a university tab needs a row added here too. Template example rows are excluded.')
print(c1,n1); print(c2,n2)
# index update
idx=wb['Index']
for row in idx.iter_rows(min_row=5):
    for c in row: c.value=None; c.hyperlink=None
rows=[('VIC','VIC Master','Built here: on-premise roll-up'),('VIC','VIC Master (Vending)','Built here: vending roll-up')]
src={'Deakin':'Deakin','La Trobe':'La Trobe','Melb Uni':'Uni_Melb','CLAYTON MONASH':'Monash','CAULFIELD MONASH':'Monash','RMIT':'RMIT','VIC UNI':'Swinburne_v1','Swinburne':'Swinburne_v1','{University':'Blank_File_v1','Calendar':'Uni_Melb / RMIT'}
for s in wb.sheetnames:
    if s in('Index','VIC Master','VIC Master (Vending)','Visit Log','VIC Progress'): continue
    f=next((v for k,v in src.items() if s.startswith(k)),None)
    if f is None: continue
    uni={'Uni_Melb':'Melb Uni','Swinburne_v1':'VU' if s.startswith('VIC UNI') else 'Swinburne','Blank_File_v1':'Template','Uni_Melb / RMIT':'All'}.get(f,f)
    rows.append(('VIC' if uni not in('Template','All') else '-',s,f'Education_Channel_Execution_Tracker_{f}.xlsx'))
idx.cell(4,1,'State'); idx.cell(4,2,'Tab'); idx.cell(4,3,'Source')
for r,(a,b,c) in enumerate(rows,5):
    idx.cell(r,1,a).font=Font(name='Arial'); x=idx.cell(r,2,b); x.hyperlink=f"#'{b}'!A1"; x.font=Font(name='Arial',color='0563C1',underline='single',bold=b.startswith('VIC Master'))
    idx.cell(r,3,c).font=Font(name='Arial',color='666666')
idx.cell(r+2,1,'The Uni Melb and RMIT files each had an identical Calendar tab; it appears once here.').font=Font(name='Arial',italic=True,color='666666')
idx.cell(r+3,1,'Other states: copy the VIC Master layout when their university tabs are added.').font=Font(name='Arial',italic=True,color='666666')
wb.active=0
for w in wb.worksheets: w.sheet_view.tabSelected=(w.title=='Index')
wb['VIC Master'].sheet_properties.tabColor=NAVY; wb['VIC Master (Vending)'].sheet_properties.tabColor=NAVY
wb.save(p)
