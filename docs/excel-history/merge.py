import sys, copy, openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
U='/root/.claude/uploads/5838e628-4997-5fab-a7b9-5672e110fb0d/'
order=[('548be836-Education_Channel_Execution_Tracker_Deakin.xlsx','Deakin'),
('1a6e1eed-Education_Channel_Execution_Tracker_La_Trobe.xlsx','La Trobe'),
('84985a1f-Education_Channel_Execution_Tracker_Uni_Melb.xlsx','Uni Melb'),
('6cfe27e3-Education_Channel_Execution_Tracker_Monash.xlsx','Monash'),
('992c66ee-Education_Channel_Execution_Tracker_RMIT.xlsx','RMIT'),
('20079ba0-Education_Channel_Execution_Tracker_Swinburne_v1.xlsx','Swinburne / VU'),
('d60da191-Education_Channel_Execution_Tracker_Blank_File_v1.xlsx','Blank template')]
out=openpyxl.Workbook(); idx=out.active; idx.title='Index'
rows=[]; seen=set()
def copy_sheet(src, dst):
    for row in src.iter_rows():
        for c in row:
            if type(c).__name__=='MergedCell': continue
            d=dst.cell(row=c.row,column=c.column,value=c.value)
            if c.has_style:
                d.font=copy.copy(c.font); d.fill=copy.copy(c.fill); d.border=copy.copy(c.border)
                d.alignment=copy.copy(c.alignment); d.number_format=c.number_format; d.protection=copy.copy(c.protection)
            if c.hyperlink: d.hyperlink=copy.copy(c.hyperlink)
            if c.comment: d.comment=copy.copy(c.comment)
    for m in src.merged_cells.ranges: dst.merge_cells(str(m))
    for k,dim in src.column_dimensions.items():
        n=dst.column_dimensions[k]; n.width=dim.width; n.hidden=dim.hidden; n.min=dim.min; n.max=dim.max; n.outlineLevel=dim.outlineLevel
    for k,dim in src.row_dimensions.items():
        n=dst.row_dimensions[k]; n.height=dim.height; n.hidden=dim.hidden; n.outlineLevel=dim.outlineLevel
    for dv in src.data_validations.dataValidation: dst.add_data_validation(copy.copy(dv))
    for rng in src.conditional_formatting:
        for r in rng.rules: dst.conditional_formatting.add(str(rng.sqref), copy.copy(r))
    dst.freeze_panes=src.freeze_panes
    dst.sheet_properties.tabColor=src.sheet_properties.tabColor
    dst.sheet_format=copy.copy(src.sheet_format)
    dst.sheet_view.showGridLines=src.sheet_view.showGridLines
    dst.sheet_view.zoomScale=src.sheet_view.zoomScale
    dst.page_setup.orientation=src.page_setup.orientation
    if src.auto_filter.ref: dst.auto_filter.ref=src.auto_filter.ref
for f,label in order:
    wb=openpyxl.load_workbook(U+f)
    for ws in wb.worksheets:
        if ws.title=='Calendar':
            if 'Calendar' in seen: continue
        name=ws.title; seen.add(name)
        dst=out.create_sheet(name); copy_sheet(ws,dst)
        rows.append((label,name,f.split('-',1)[1]))
# index
idx['A1']='Education Channel Execution Tracker — combined'; idx['A1'].font=Font(name='Arial',bold=True,size=14)
idx['A2']='All university trackers in one workbook. Click a tab name to jump to it.'; idx['A2'].font=Font(name='Arial',italic=True,color='666666')
hdr=['University','Tab','Source file']
for i,h in enumerate(hdr,1):
    c=idx.cell(row=4,column=i,value=h); c.font=Font(name='Arial',bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3864')
for r,(label,name,src) in enumerate(rows,5):
    idx.cell(row=r,column=1,value=label).font=Font(name='Arial')
    c=idx.cell(row=r,column=2,value=name); c.hyperlink=f"#'{name}'!A1"; c.font=Font(name='Arial',color='0563C1',underline='single')
    idx.cell(row=r,column=3,value=src).font=Font(name='Arial',color='666666')
idx.cell(row=r+2,column=1,value='Note: the Uni Melb and RMIT files each had an identical Calendar tab; it appears once here.').font=Font(name='Arial',italic=True,color='666666')
idx.column_dimensions['A'].width=20; idx.column_dimensions['B'].width=34; idx.column_dimensions['C'].width=62
out.save(sys.argv[1]); print(len(rows),'tabs + Index')
